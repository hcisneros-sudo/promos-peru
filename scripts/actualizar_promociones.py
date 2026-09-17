from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE_FILES = {
    "promociones_interbank.xlsx": "Interbank",
    "BanBif_ClubHola_Promociones.xlsx": "BanBif",
    "Cencosud_Promociones.xlsx": "Tarjeta Cencosud",
    "SIP_Promociones.xlsx": "SIP",
    "promociones_banco_falabella.xlsx": "Banco Falabella",
    "promociones_banco_ripley.xlsx": "Banco Ripley",
}

CATEGORY_RULES = [
    ("Gastronomía", [
        ("Cafeterías y pastelerías", ["cafe", "cafeteria", "pasteleria", "starbucks", "dunkin"]),
        ("Comida rápida", ["fast food", "hamburgues", "pizza", "pollo", "polleria", "burger", "kfc", "bembos", "papas"]),
        ("Bares y bebidas", ["bar ", "cerveza", "licor", "coctel", "bebidas"]),
        ("Restaurantes", ["restaur", "cevicher", "gastronom", "comida", "chifa", "sushi", "parrilla", "restofans"]),
    ]),
    ("Supermercados y alimentos", [
        ("Supermercados", ["supermerc", "metro", "wong", "vivanda", "plaza vea", "tottus"]),
        ("Minimarkets", ["minimarket", "bodega", "tienda de conveniencia"]),
        ("Alimentos y bebidas", ["alimentos", "abarrotes", "productos naturales", "bodega"]),
    ]),
    ("Moda y accesorios", [
        ("Calzado", ["calzado", "zapato", "zapatilla", "sneaker"]),
        ("Accesorios y joyería", ["joyeria", "reloj", "accesorio", "cartera", "lentes de sol"]),
        ("Ropa", ["moda", "ropa", "terno", "camisa", "pantalon", "vestido", "textil"]),
    ]),
    ("Salud y bienestar", [
        ("Farmacias", ["farmacia", "botica", "inkafarma", "mifarma"]),
        ("Clínicas y laboratorios", ["clinica", "laboratorio", "medico", "dental", "odontolog"]),
        ("Ópticas", ["optica", "oftalmolog"]),
        ("Gimnasios", ["gimnasio", "fitness", "entrenamiento", "smart fit"]),
        ("Belleza y cuidado personal", ["belleza", "salon", "spa", "estetica", "peluquer", "barber", "montalvo"]),
    ]),
    ("Viajes y transporte", [
        ("Vuelos", ["aerolinea", "vuelo", "latam", "sky airline", "jetsmart"]),
        ("Hoteles", ["hotel", "alojamiento", "hospedaje", "resort"]),
        ("Turismo", ["turismo", "tour", "viaje", "paquete turistico"]),
        ("Movilidad", ["taxi", "movilidad", "transporte", "cabify", "uber"]),
        ("Combustible y automotriz", ["automotr", "combustible", "grifo", "gasolina", "mantenimiento", "lavado", "llantas"]),
    ]),
    ("Entretenimiento", [
        ("Cine", ["cine", "cineplanet", "cinemark", "movie time"]),
        ("Espectáculos", ["concierto", "teatro", "entrada", "espectaculo", "evento"]),
        ("Parques y experiencias", ["parque", "experiencia", "bungee", "escape room", "aventura"]),
        ("Juegos", ["gaming", "videojuego", "juegos", "bowling"]),
    ]),
    ("Tecnología y electro", [
        ("Celulares y computación", ["celular", "smartphone", "laptop", "computador", "tablet", "iphone"]),
        ("Electrodomésticos", ["electrodomest", "refrigeradora", "televisor", "lavadora"]),
        ("Tecnología", ["tecnolog", "electronica", "software", "audio"]),
    ]),
    ("Hogar", [
        ("Muebles y decoración", ["mueble", "decoracion", "colchon", "hogar"]),
        ("Ferretería y mejoramiento", ["ferreter", "mejoramiento", "construccion", "pintura"]),
        ("Servicios para el hogar", ["limpieza del hogar", "servicio domestico", "reparacion del hogar"]),
    ]),
    ("Educación", [
        ("Instituciones educativas", ["universidad", "instituto", "colegio", "escuela"]),
        ("Idiomas", ["idioma", "ingles", "britanico", "icpna"]),
        ("Cursos", ["curso", "capacitacion", "diplomado", "taller"]),
    ]),
    ("Mascotas", [
        ("Veterinarias", ["veterinar", "clinica veterinaria"]),
        ("Alimentos y accesorios", ["mascota", "pet shop", "superpet", "alimento para perro", "alimento para gato"]),
    ]),
    ("Compras", [
        ("Tiendas por departamento", ["tienda por departamento", "falabella", "ripley", "oechsle"]),
        ("Centros comerciales", ["centro comercial", "mall ", "real plaza", "jockey plaza", "plaza san miguel"]),
        ("Comercio electrónico", ["marketplace", "e-commerce", "comercio electronico", "compra online"]),
    ]),
    ("Servicios", [
        ("Seguros y finanzas", ["seguro", "prestamo", "credito", "cuenta sueldo", "financ", "tasa"]),
        ("Telecomunicaciones", ["telefonia", "internet", "movistar", "claro", "entel", "bitel"]),
        ("Aplicaciones y delivery", ["delivery", "aplicacion", "app ", "rappi", "pedidosya"]),
        ("Suscripciones", ["suscripcion", "membresia", "streaming"]),
        ("Otros servicios", ["servicio", "consultoria", "fotografia", "imprenta"]),
    ]),
]

MERCHANT_ALIASES = {
    "starbucks coffee": "Starbucks", "starbucks peru": "Starbucks",
    "cine planet": "Cineplanet", "cineplanet peru": "Cineplanet",
    "inkafarma boticas": "Inkafarma", "boticas mifarma": "Mifarma",
    "montalvo salon spa": "Montalvo Salón & Spa",
}

DISTRICTS = [
    "Ancón", "Ate", "Barranco", "Bellavista", "Breña", "Callao", "Carabayllo",
    "Carmen de la Legua", "Cercado de Lima", "Chaclacayo", "Chorrillos", "Cieneguilla",
    "Comas", "El Agustino", "Independencia", "Jesús María", "La Molina", "La Perla",
    "La Victoria", "Lince", "Los Olivos", "Lurín", "Magdalena del Mar", "Miraflores",
    "Pueblo Libre", "Puente Piedra", "Rímac", "San Borja", "San Isidro",
    "San Juan de Lurigancho", "San Juan de Miraflores", "San Luis",
    "San Martín de Porres", "San Miguel", "Santa Anita", "Santiago de Surco",
    "Surquillo", "Villa El Salvador", "Villa María del Triunfo",
    "La Punta", "Mi Perú", "Ventanilla",
]

CALLAO_DISTRICTS = {"Bellavista", "Callao", "Carmen de la Legua", "La Perla", "La Punta", "Mi Perú", "Ventanilla"}

DISTRICT_ALIASES = {
    "surco": "Santiago de Surco", "magdalena": "Magdalena del Mar",
    "san martin de porres": "San Martín de Porres", "jesus maria": "Jesús María",
    "rimac": "Rímac", "san juan de miraflores": "San Juan de Miraflores",
    "sjm": "San Juan de Miraflores", "sjl": "San Juan de Lurigancho",
}

# Localidades que aparecieron dentro del texto de la dirección, pero no en las
# columnas estructuradas de los Excel. Se evalúan antes que los distritos de
# Lima para evitar errores como interpretar "Independencia, Huánuco" como Lima.
PERU_LOCALITIES = [
    (("juliaca",), "Juliaca", "San Román", "Puno"),
    (("chiclayo",), "Chiclayo", "Chiclayo", "Lambayeque"),
    (("trujillo",), "Trujillo", "Trujillo", "La Libertad"),
    (("huancayo",), "Huancayo", "Huancayo", "Junín"),
    (("ayacucho", "huamanga"), "Ayacucho", "Huamanga", "Ayacucho"),
    (("cajamarca",), "Cajamarca", "Cajamarca", "Cajamarca"),
    (("arequipa",), "Arequipa", "Arequipa", "Arequipa"),
    (("chimbote",), "Chimbote", "Santa", "Áncash"),
    (("huaraz",), "Huaraz", "Huaraz", "Áncash"),
    (("huanuco",), "Huánuco", "Huánuco", "Huánuco"),
    (("tarapoto",), "Tarapoto", "San Martín", "San Martín"),
    (("moyobamba",), "Moyobamba", "Moyobamba", "San Martín"),
    (("pucallpa",), "Pucallpa", "Coronel Portillo", "Ucayali"),
    (("iquitos",), "Iquitos", "Maynas", "Loreto"),
    (("puerto maldonado",), "Puerto Maldonado", "Tambopata", "Madre de Dios"),
    (("cerro de pasco",), "Cerro de Pasco", "Pasco", "Pasco"),
    (("chachapoyas",), "Chachapoyas", "Chachapoyas", "Amazonas"),
    (("abancay",), "Abancay", "Abancay", "Apurímac"),
    (("cusco", "cuzco"), "Cusco", "Cusco", "Cusco"),
    (("piura",), "Piura", "Piura", "Piura"),
    (("ica",), "Ica", "Ica", "Ica"),
    (("ilo",), "Ilo", "Ilo", "Moquegua"),
    (("moquegua",), "Moquegua", "Mariscal Nieto", "Moquegua"),
    (("tacna",), "Tacna", "Tacna", "Tacna"),
    (("tumbes",), "Tumbes", "Tumbes", "Tumbes"),
    (("puno",), "Puno", "Puno", "Puno"),
]

DEPARTMENT_ALIASES = {
    "amazonas": "Amazonas", "ancash": "Áncash", "apurimac": "Apurímac",
    "arequipa": "Arequipa", "ayacucho": "Ayacucho", "cajamarca": "Cajamarca",
    "cusco": "Cusco", "huancavelica": "Huancavelica", "huanuco": "Huánuco",
    "ica": "Ica", "junin": "Junín", "la libertad": "La Libertad",
    "lambayeque": "Lambayeque", "loreto": "Loreto", "madre de dios": "Madre de Dios",
    "moquegua": "Moquegua", "pasco": "Pasco", "piura": "Piura", "puno": "Puno",
    "san martin": "San Martín", "tacna": "Tacna", "tumbes": "Tumbes", "ucayali": "Ucayali",
}

NATIONAL_TERMS = ["nivel nacional", "todo el peru", "a nivel nacional", "todos los locales"]
ONLINE_TERMS = ["solo online", "exclusivo online", "pagina web", "aplicativo", "e-commerce", "web y app"]
PERU_BOUNDS = {"lat_min": -18.6, "lat_max": 0.6, "lon_min": -82.0, "lon_max": -68.0}


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=" ")
    text = re.sub(r"\s+", " ", str(value)).strip()
    if text.lower() in {"none", "nan", "null", "--", "no disponible"} or " at 0x" in text:
        return ""
    return text


def clean_multiline(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    if text.lower() in {"none", "nan", "null", "--", "no disponible"} or " at 0x" in text:
        return ""
    return text


def fold(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", value.lower()) if not unicodedata.combining(c))


def key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", fold(value)).strip()


def first(row: dict, *columns: str) -> str:
    for column in columns:
        value = clean(row.get(column))
        if value:
            return value
    return ""


def first_multiline(row: dict, *columns: str) -> str:
    for column in columns:
        value = clean_multiline(row.get(column))
        if value:
            return value
    return ""


def parse_date(value: str) -> str:
    value = clean(value)
    if not value:
        return ""
    for pattern in (r"(20\d{2})[-/](\d{1,2})[-/](\d{1,2})", r"(\d{1,2})[-/](\d{1,2})[-/](20\d{2})"):
        match = re.search(pattern, value)
        if not match:
            continue
        if pattern.startswith("(20"):
            year, month, day = match.groups()
        else:
            day, month, year = match.groups()
        try:
            return date(int(year), int(month), int(day)).isoformat()
        except ValueError:
            return ""
    return ""


def normalize_merchant(value: str) -> str:
    value = clean(value)
    basic = re.sub(r"\b(s\.?a\.?c\.?|s\.?a\.?|e\.?i\.?r\.?l\.?)\b", "", value, flags=re.I)
    basic = re.sub(r"\s+", " ", basic).strip(" -–—|,")
    return MERCHANT_ALIASES.get(key(basic), basic)


def classify(source_category: str, merchant: str, name: str, detail: str) -> tuple[str, str, str, str]:
    parts = [(merchant, 5), (name, 4), (source_category, 3), (detail, 1)]
    scores = defaultdict(int)
    hits = defaultdict(list)
    for text, weight in parts:
        haystack = fold(text)
        for category, subcategories in CATEGORY_RULES:
            for subcategory, terms in subcategories:
                matched = [term for term in terms if term in haystack]
                if matched:
                    scores[(category, subcategory)] += weight * len(matched)
                    hits[(category, subcategory)].extend(matched)
    if not scores:
        return "Otros", "Sin clasificar", "Revisar", "Sin coincidencias"
    ordered = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    (category, subcategory), score = ordered[0]
    tied = len(ordered) > 1 and ordered[1][1] == score
    confidence = "Revisar" if tied else ("Alta" if score >= 5 else "Media" if score >= 3 else "Baja")
    evidence = ", ".join(sorted(set(hits[(category, subcategory)])))[:200]
    return category, subcategory, confidence, evidence


def benefit_type(text: str) -> str:
    value = fold(text)
    if "cashback" in value or "devolucion" in value:
        return "Cashback"
    if re.search(r"\b2\s*[x×]\s*1\b", value):
        return "2x1"
    if "cuota" in value and ("sin interes" in value or "0%" in value):
        return "Cuotas sin intereses"
    if "%" in value or "descuento" in value or "dscto" in value or "dto" in value:
        return "Descuento"
    if "precio" in value or re.search(r"\bs/\s*\d", value):
        return "Precio especial"
    return "Otro beneficio"


def extract_discount(explicit: str, text: str):
    candidates = []
    for value in (explicit, text):
        for match in re.findall(r"(?<!\d)(\d{1,3}(?:[.,]\d+)?)\s*%", clean(value)):
            number = float(match.replace(",", "."))
            if 0 <= number <= 100:
                candidates.append(number)
    return max(candidates) if candidates else None


def calculate_status(start_date: str, end_date: str) -> str:
    today = date.today()
    start = date.fromisoformat(start_date) if start_date else None
    end = date.fromisoformat(end_date) if end_date else None
    if start and today < start:
        return "Próxima"
    if end and today > end:
        return "Vencida"
    if start or end:
        return "Vigente"
    return "Sin fecha"


def detect_scope(location: str, modality: str) -> str:
    value = fold(f"{location} {modality}")
    if any(term in value for term in ONLINE_TERMS):
        return "Online"
    if any(term in value for term in NATIONAL_TERMS):
        return "Nacional"
    return "Local" if location else "Sin determinar"


def detect_district(text: str) -> str:
    value = fold(text)
    for alias, canonical in sorted(DISTRICT_ALIASES.items(), key=lambda item: -len(item[0])):
        if re.search(rf"\b{re.escape(alias)}\b", value):
            return canonical
    for district in sorted(DISTRICTS, key=len, reverse=True):
        if re.search(rf"\b{re.escape(fold(district))}\b", value):
            return district
    return ""


def detect_geography(text: str) -> tuple[str, str, str]:
    """Devuelve distrito/ciudad, provincia y departamento sin asumir Lima."""
    value = fold(text)
    for aliases, city, province, department in PERU_LOCALITIES:
        if any(re.search(rf"\b{re.escape(alias)}\b", value) for alias in aliases):
            return city, province, department

    # Un departamento explícito es mejor que completar erróneamente con Lima.
    for alias, department in sorted(DEPARTMENT_ALIASES.items(), key=lambda item: -len(item[0])):
        if re.search(rf"\b{re.escape(alias)}\b", value):
            return "", "", department

    district = detect_district(text)
    if district:
        region = "Callao" if district in CALLAO_DISTRICTS else "Lima"
        return district, region, region
    return "", "", ""


def normalize_address(text: str) -> str:
    text = unicodedata.normalize("NFKC", clean(text))
    text = re.sub(r"[\u200b-\u200f\u2060\ufeff]", "", text)
    text = text.replace("“", '"').replace("”", '"').replace("’", "'")
    text = re.sub(r"\b(?:avenida\b|av(?=\s*\.|\s))\s*\.*\s*", "Av. ", text, flags=re.I)
    text = re.sub(r"\b(?:jir[oó]n\b|jr(?=\s*\.|\s))\s*\.*\s*", "Jr. ", text, flags=re.I)
    text = re.sub(r"\b(?:carretera\b|carret(?=\s*\.|\s))\s*\.*\s*", "Carretera ", text, flags=re.I)
    text = re.sub(r"\bcalle\b\s*\.*\s*", "Calle ", text, flags=re.I)
    text = re.sub(r"\b(?:n(?:ro|um(?:ero)?)?(?:\s*[.°º])+|nro\.?|#)\s*(?=\d)", "", text, flags=re.I)
    text = re.sub(r"\.{2,}", ".", text)
    text = re.sub(r"\s*([,:;])\s*", r"\1 ", text)
    text = re.sub(r"[\"'](?=\s|$)", "", text)
    text = re.sub(r"\bSurcol\b", "Surco", text, flags=re.I)
    text = re.sub(r"\bJes[uú]s Mar[ií]a[\"']?l\b", "Jesús María", text, flags=re.I)
    text = re.sub(r"\bProlog(?:aci[oó]n|\.)?\b", "Prolongación", text, flags=re.I)
    text = re.sub(r"\bMall\s+Av\.\s*entura\b", "Mall Aventura", text, flags=re.I)
    text = re.sub(r"\bLambayaque\b", "Lambayeque", text, flags=re.I)
    return re.sub(r"\s+", " ", text).strip(" -–—|,;:")


def parse_coordinate_number(raw: str, axis: str) -> float | None:
    """Repara coma decimal, signo perdido y decimales omitidos con validación Perú."""
    value = raw.strip().replace(" ", "").replace(",", ".")
    try:
        number = float(value)
    except ValueError:
        return None
    limit = 90 if axis == "lat" else 180
    while abs(number) > limit:
        number /= 10
    if axis == "lat" and number > PERU_BOUNDS["lat_max"]:
        number = -number
    if axis == "lon" and number > 0:
        number = -number
    return number


def extract_coordinates(text: str) -> tuple[float | None, float | None, str]:
    pattern = re.compile(r"\(\s*([+-]?\d{1,10}(?:[.,]\d+)?)\s*[,;]\s*([+-]?\d{1,11}(?:[.,]\d+)?)\s*\)")
    for match in pattern.finditer(text):
        lat = parse_coordinate_number(match.group(1), "lat")
        lon = parse_coordinate_number(match.group(2), "lon")
        if lat is None or lon is None:
            continue
        if PERU_BOUNDS["lat_min"] <= lat <= PERU_BOUNDS["lat_max"] and PERU_BOUNDS["lon_min"] <= lon <= PERU_BOUNDS["lon_max"]:
            cleaned = f"{text[:match.start()]} {text[match.end():]}"
            return lat, lon, re.sub(r"\s+", " ", cleaned).strip()
    return None, None, text


def looks_like_address(text: str) -> bool:
    value = fold(text)
    has_route = bool(re.search(r"\b(av|avenida|jr|jiron|calle|carretera|prolongacion|pasaje|malecon)\b", value))
    return has_route or bool(re.search(r"\d{2,}", value))


def separate_place_and_address(text: str) -> tuple[str, str]:
    """Separa 'nombre del local: dirección' sin perder el valor original."""
    cleaned = normalize_address(text)
    for separator in (":", " — ", " – "):
        if separator not in cleaned:
            continue
        left, right = cleaned.split(separator, 1)
        if looks_like_address(right):
            return left.strip(" ,:-"), normalize_address(right)

    # Algunos bancos concatenan el centro comercial y la vía sin separador.
    route = re.search(r"\b(?:Av\.|Jr\.|Calle|Carretera|Prolongaci[oó]n|Pasaje|Malec[oó]n)\s+", cleaned, flags=re.I)
    if route and route.start() > 0 and re.search(r"\d", cleaned[route.start():]):
        prefix = cleaned[:route.start()].strip(" ,:-")
        if re.search(r"\b(?:c\.?\s*c\.?|centro comercial|mall|real plaza|plaza|tienda|local)\b", fold(prefix)):
            return prefix, normalize_address(cleaned[route.start():])
    return "", cleaned


def split_locations(raw: str) -> list[str]:
    raw = clean_multiline(raw)
    if not raw or any(term in fold(raw) for term in NATIONAL_TERMS):
        return []
    if len(raw) < 90 and not re.search(r"\d", raw) and detect_district(raw):
        return []
    pieces = re.split(r"\s*\|\s*|\s*;\s*|\s*•\s*|\n+|(?=\b(?:LOCAL|SEDE|TIENDA)\s+\d+\b)", raw, flags=re.I)
    results = []
    for piece in pieces:
        piece = clean(piece).lstrip("• ")
        if len(piece) < 8 or piece in results:
            continue
        results.append(piece)
    return results


def location_records(promotion_id: str, location_text: str, source: str) -> list[dict]:
    rows = []
    seen = set()
    for index, original in enumerate(split_locations(location_text), 1):
        latitude, longitude, without_coords = extract_coordinates(original)
        name, address = separate_place_and_address(without_coords)
        normalized_key = (key(name), key(address))
        if normalized_key in seen:
            continue
        seen.add(normalized_key)
        district, province, department = detect_geography(f"{name} {address}")
        precision = "Exacta" if latitude is not None else ("Dirección por geocodificar" if re.search(r"\d", address) else "Revisar")
        has_number = bool(re.search(r"\d", address))
        looks_searchable = looks_like_address(address) or bool(re.search(r"\b(?:mall|plaza|centro comercial|c\.?\s*c\.?)\b", fold(f"{name} {address}")))
        confidence = "Alta" if latitude is not None or ((district or department) and has_number) else "Media" if looks_searchable else "Revisar"
        local_id = hashlib.sha1(f"{promotion_id}|{key(address)}|{key(name)}".encode()).hexdigest()[:18]
        rows.append({
            "id_local": local_id, "id_promocion": promotion_id, "nombre_local": name,
            "direccion_original": original, "direccion_limpia": address, "distrito": district,
            "provincia": province, "departamento": department,
            "latitud": latitude, "longitud": longitude,
            "precision_ubicacion": precision, "fuente_coordenada": "Excel original normalizado" if latitude is not None else "Pendiente de geocodificación",
            "confianza_ubicacion": confidence, "archivo_origen": source,
        })
    return rows


def source_mapping(row: dict, filename: str) -> dict:
    bank = first(row, "Banco", "Banco / marca") or SOURCE_FILES[filename]
    name = first(row, "Nombre", "Promoción", "Establecimiento")
    merchant = normalize_merchant(first(row, "Comercio", "Establecimiento") or name)
    detail = first(row, "Detalle de la promoción", "Descripción", "Resumen", "Promoción")
    benefit = first(row, "Texto del beneficio", "Beneficio", "Precio o beneficio adicional", "Descuento (%)")
    source_category = first(row, "Categoría", "Categorías", "Subcategorías", "Tipo", "Categoría ID")
    location = first_multiline(row, "Sucursales") or first_multiline(row, "Locales") or first_multiline(row, "Ubicación")
    regional = first_multiline(row, "Regiones / zonas", "Regiones", "Ciudades")
    if not location:
        location = regional
    modality = first(row, "Modalidad")
    start_date = parse_date(first(row, "Fecha de inicio", "Fecha inicio"))
    end_date = parse_date(first(row, "Fecha de fin", "Fecha de término", "Fecha fin"))
    source_id = first(row, "ID promoción", "ID", "ID CMS")
    url = first(row, "Link", "URL de la promoción", "URL", "Enlace externo")
    stable = f"{bank}|{source_id or url or name}|{merchant}"
    promotion_id = hashlib.sha1(stable.encode()).hexdigest()[:18]
    category, subcategory, category_confidence, category_evidence = classify(source_category, merchant, name, detail)
    benefit_text = benefit or detail or name
    return {
        "id_promocion": promotion_id, "id_fuente": source_id, "banco": bank,
        "comercio": merchant, "nombre_promocion": name or merchant or "Promoción sin título",
        "descripcion": detail, "categoria": category, "subcategoria": subcategory,
        "confianza_categoria": category_confidence, "evidencia_categoria": category_evidence,
        "categoria_original": source_category, "tipo_beneficio": benefit_type(benefit_text),
        "porcentaje_descuento": extract_discount(first(row, "Descuento (%)"), benefit_text),
        "beneficio_texto": benefit_text, "medio_pago": first(row, "Productos / medios de pago", "Tarjetas válidas", "Tarjeta / medio de pago", "Niveles ClubHOLA"),
        "modalidad": modality, "alcance": detect_scope(f"{location} {regional}", modality),
        "fecha_inicio": start_date, "fecha_fin": end_date,
        "dias_validos": first(row, "Días válidos"), "estado_vigencia": calculate_status(start_date, end_date),
        "estado_original": first(row, "Estado de vigencia", "Estado"),
        "restricciones": first(row, "Restricciones", "Restricciones / vigencia", "Disponibilidad"),
        "terminos_condiciones": first(row, "Términos y condiciones"),
        "ubicacion_original": location, "regiones_originales": regional,
        "url_promocion": url, "url_imagen": first(row, "Imagen", "Imagen web", "Imagen grande", "Miniatura"),
        "archivo_origen": filename, "fecha_extraccion": first(row, "Fecha de extracción (UTC)", "Fecha extracción UTC", "Fecha de extracción"),
    }


def propagate_location_data(locations: list[dict]) -> None:
    """Completa geografía/coordenadas cuando la misma dirección ya las aporta en otra fila."""
    groups = defaultdict(list)
    for location in locations:
        signature = key(location.get("direccion_limpia", ""))
        if len(signature) >= 6:
            groups[signature].append(location)

    for rows in groups.values():
        geography = {
            (row.get("distrito", ""), row.get("provincia", ""), row.get("departamento", ""))
            for row in rows if row.get("distrito") or row.get("provincia") or row.get("departamento")
        }
        if len(geography) == 1:
            district, province, department = next(iter(geography))
            for row in rows:
                row["distrito"] = row.get("distrito") or district
                row["provincia"] = row.get("provincia") or province
                row["departamento"] = row.get("departamento") or department

        coordinates = {
            (round(float(row["latitud"]), 7), round(float(row["longitud"]), 7))
            for row in rows if row.get("latitud") is not None and row.get("longitud") is not None
        }
        if len(coordinates) == 1 and re.search(r"\d", rows[0].get("direccion_limpia", "")):
            latitude, longitude = next(iter(coordinates))
            for row in rows:
                if row.get("latitud") is None:
                    row["latitud"], row["longitud"] = latitude, longitude
                    row["precision_ubicacion"] = "Exacta por dirección repetida"
                    row["fuente_coordenada"] = "Reutilizada del Excel original"
                    row["confianza_ubicacion"] = "Alta"


def add_quality_issues(promo: dict, locations: list[dict], issues: list[dict]):
    def issue(kind: str, field: str, found: str, reason: str, suggestion: str):
        issues.append({"id_promocion": promo["id_promocion"], "banco": promo["banco"], "comercio": promo["comercio"], "tipo_observacion": kind, "campo": field, "valor_encontrado": clean(found), "motivo": reason, "correccion_sugerida": suggestion, "accion_realizada": "Pendiente"})
    if not promo["comercio"]:
        issue("Dato obligatorio", "comercio", "", "No se identificó el comercio.", "Revisar nombre y descripción.")
    if promo["confianza_categoria"] in {"Baja", "Revisar"}:
        issue("Categoría", "categoria", promo["categoria"], f"Clasificación con confianza {promo['confianza_categoria'].lower()}.", "Seleccionar categoría y subcategoría del catálogo maestro.")
    if promo["porcentaje_descuento"] is not None and not 0 <= promo["porcentaje_descuento"] <= 100:
        issue("Beneficio", "porcentaje_descuento", promo["porcentaje_descuento"], "Porcentaje fuera del rango permitido.", "Verificar el texto del beneficio.")
    if promo["fecha_inicio"] and promo["fecha_fin"] and promo["fecha_inicio"] > promo["fecha_fin"]:
        issue("Fechas", "fecha_fin", promo["fecha_fin"], "La fecha final es anterior a la inicial.", "Revisar vigencia publicada.")
    if promo["alcance"] == "Local" and promo["ubicacion_original"] and not locations:
        issue("Dirección", "ubicacion_original", promo["ubicacion_original"], "No se pudo separar una dirección utilizable.", "Revisar si corresponde a región, distrito o dirección.")
    for loc in locations:
        if loc["confianza_ubicacion"] == "Revisar":
            issue("Dirección", "direccion_original", loc["direccion_original"], "Dirección ambigua o incompleta.", "Completar dirección o identificar el local.")


def read_sources(input_dir: Path):
    promotions, locations, issues = [], [], []
    missing = []
    for filename in SOURCE_FILES:
        path = input_dir / filename
        if not path.exists():
            missing.append(filename)
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["Promociones"] if "Promociones" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows)]
        for values in rows:
            row = dict(zip(headers, values))
            if not any(clean(value) for value in values):
                continue
            promo = source_mapping(row, filename)
            promo_locations = location_records(promo["id_promocion"], promo["ubicacion_original"], filename)
            promotions.append(promo)
            locations.extend(promo_locations)
        workbook.close()
    propagate_location_data(locations)
    locations_by_promotion = defaultdict(list)
    for location in locations:
        locations_by_promotion[location["id_promocion"]].append(location)
    for promo in promotions:
        add_quality_issues(promo, locations_by_promotion[promo["id_promocion"]], issues)
    return promotions, locations, issues, missing


def detect_duplicates(promotions: list[dict], issues: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for promo in promotions:
        signature = "|".join([key(promo["banco"]), key(promo["comercio"]), key(promo["nombre_promocion"]), str(promo["porcentaje_descuento"] or ""), promo["fecha_inicio"], promo["fecha_fin"]])
        groups[signature].append(promo)
    duplicates = []
    for signature, rows in groups.items():
        if len(rows) < 2:
            continue
        group_id = hashlib.sha1(signature.encode()).hexdigest()[:12]
        for promo in rows:
            duplicates.append({"grupo_duplicado": group_id, "id_promocion": promo["id_promocion"], "banco": promo["banco"], "comercio": promo["comercio"], "nombre_promocion": promo["nombre_promocion"], "beneficio_texto": promo["beneficio_texto"], "fecha_inicio": promo["fecha_inicio"], "fecha_fin": promo["fecha_fin"], "url_promocion": promo["url_promocion"], "archivo_origen": promo["archivo_origen"]})
    return duplicates


def build_summary(promotions, locations, issues, duplicates, missing):
    by_bank = Counter(row["banco"] for row in promotions)
    by_category = Counter(row["categoria"] for row in promotions)
    return {
        "fecha_procesamiento": datetime.now().isoformat(timespec="seconds"),
        "archivos_faltantes": missing, "total_promociones": len(promotions),
        "total_locales_separados": len(locations),
        "locales_con_coordenadas": sum(row["latitud"] is not None for row in locations),
        "locales_pendientes_geocodificar": sum(row["latitud"] is None and row["confianza_ubicacion"] != "Revisar" for row in locations),
        "observaciones_calidad": len(issues), "filas_duplicados_probables": len(duplicates),
        "promociones_por_banco": dict(by_bank), "promociones_por_categoria": dict(by_category),
    }


def main():
    parser = argparse.ArgumentParser(description="Consolida y limpia promociones bancarias.")
    parser.add_argument("--entrada", default="fuentes", help="Carpeta con los seis Excel originales")
    parser.add_argument("--salida", default="salida", help="Carpeta para los resultados JSON")
    args = parser.parse_args()
    input_dir, output_dir = Path(args.entrada), Path(args.salida)
    output_dir.mkdir(parents=True, exist_ok=True)
    promotions, locations, issues, missing = read_sources(input_dir)
    duplicates = detect_duplicates(promotions, issues)
    summary = build_summary(promotions, locations, issues, duplicates, missing)
    datasets = {"promociones": promotions, "locales": locations, "control_calidad": issues, "duplicados_probables": duplicates, "resumen": summary}
    for name, data in datasets.items():
        (output_dir / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Promociones procesadas: {len(promotions):,}")
    print(f"Locales separados: {len(locations):,}")
    print(f"Observaciones para revisión: {len(issues):,}")
    print(f"Salida: {output_dir.resolve()}")


if __name__ == "__main__":
    main()
