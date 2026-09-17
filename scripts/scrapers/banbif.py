#!/usr/bin/env python3
"""Extrae promociones de BanBif / ClubHOLA y genera un Excel.

Modos:
  dni      : Chrome inicia sesión y la API descarga el catálogo (predeterminado).
  public   : campañas visibles sin iniciar sesión.
  token    : API autenticada; requiere CLUBHOLA_TOKEN en el entorno.

Instalación:
  pip install selenium pandas openpyxl

Ejemplos:
  python BanBif_ClubHola_Promociones_API.py
  python BanBif_ClubHola_Promociones_API.py --mostrar-navegador
  python BanBif_ClubHola_Promociones_API.py --modo dni --categoria gastronomia
  python BanBif_ClubHola_Promociones_API.py --modo dni --sin-detalle
  python BanBif_ClubHola_Promociones_API.py --modo public --solo-vigentes
"""

from __future__ import annotations

import argparse
import base64
import html
import os
import re
import sys
import time
import unicodedata
import json
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HOME = "https://www.clubhola.pe"
API = f"{HOME}/apimovil"
DEFAULT_OUTPUT = "BanBif_ClubHola_Promociones.xlsx"
DEFAULT_DNI = ""
TIMEOUT = 45

HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json",
    "Origin": HOME,
    "Referer": f"{HOME}/",
    "Sec-Fetch-Dest": "empty",
    "Sec-Fetch-Mode": "cors",
    "Sec-Fetch-Site": "same-origin",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 Chrome/140.0.0.0 Safari/537.36"
    ),
}

COLUMNS = [
    "Banco", "Tipo", "Nombre", "Comercio", "Detalle de la promoción",
    "Cómo obtener el descuento", "Restricciones / vigencia", "Términos y condiciones",
    "Ubicación", "Sucursales", "Link", "Enlace online", "Web del comercio",
    "Categoría", "Subcategoría", "Campañas", "ID promoción", "ID comercio",
    "Beneficio", "Prefijo beneficio", "Sufijo beneficio", "Fecha de inicio",
    "Fecha de fin", "Estado de vigencia", "Solo por hoy", "Niveles ClubHOLA",
    "Teléfono", "Email", "Stock total", "Stock disponible", "Sin stock",
    "Requiere cupón", "Compra con login", "Máximo por suscriptor",
    "Valoración promedio", "Cantidad de valoraciones", "Imagen",
    "Fecha de actualización", "Cobertura", "Estado de extracción",
    "Fecha extracción UTC",
]


def clean(value: Any) -> str:
    """Convierte HTML/texto de la API en una celda legible."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return " | ".join(clean(x) for x in value if clean(x))
    text = html.unescape(str(value))
    text = re.sub(r"<\s*br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"</(?:p|div|li|h\d)\s*>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("\xa0", " ")
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", text)).strip()


def normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean(value).lower())
    return "".join(c for c in text if not unicodedata.combining(c))


def truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "si", "sí", "yes"}


def join_dicts(items: Any) -> str:
    if not isinstance(items, list):
        return clean(items)
    values = []
    for item in items:
        if isinstance(item, dict):
            name = clean(item.get("nombre") or item.get("name"))
            address = clean(item.get("direccion") or item.get("address"))
            values.append(" — ".join(x for x in (name, address) if x))
        else:
            values.append(clean(item))
    return " | ".join(x for x in values if x)


def promo_url(item: dict[str, Any]) -> str:
    explicit = clean(item.get("url_beneficio"))
    if explicit:
        return explicit if explicit.startswith("http") else f"{HOME}/{explicit.lstrip('/')}"
    filename = Path(urlparse(clean(item.get("path_logo"))).path).name
    route = re.sub(r"\.([^.]+)$", "", filename)
    route = re.sub(r"-\d+x\d+(?=-|$)", "", route)
    if route:
        return f"{HOME}/{route}"
    slug, promo_id = clean(item.get("slug")), clean(item.get("id"))
    return f"{HOME}/{slug}-{promo_id}" if slug and promo_id else HOME


class ClubHolaClient:
    def __init__(self, token: str = "") -> None:
        self.headers = dict(HEADERS)
        if token:
            self.headers["Authorization"] = f"Bearer {token}"

    def get(self, endpoint: str, **params: Any) -> dict[str, Any]:
        url = f"{API}/{endpoint}"
        last_error: Exception | None = None
        for attempt in range(4):
            try:
                query = urlencode({"brand": "clubhola", **params})
                request = Request(f"{url}?{query}", headers=self.headers)
                with urlopen(request, timeout=TIMEOUT) as response:
                    payload = json.load(response)
                status = payload.get("status")
                if status not in (None, 0, "0"):
                    raise RuntimeError(clean(payload.get("message")) or f"status={status}")
                return payload
            except HTTPError as exc:
                if exc.code == 401:
                    raise PermissionError(
                        f"{endpoint} requiere una sesión válida (HTTP 401)."
                    ) from exc
                last_error = exc
                if attempt == 3:
                    break
                time.sleep(1.5 * (2**attempt))
            except PermissionError:
                raise
            except (URLError, TimeoutError, ValueError, RuntimeError) as exc:
                last_error = exc
                if attempt == 3:
                    break
                time.sleep(1.5 * (2**attempt))
        raise RuntimeError(f"No se pudo consultar {endpoint}: {last_error}")


def fetch_public(client: ClubHolaClient, workers: int) -> list[dict[str, Any]]:
    concepts = client.get("getConceptos").get("data") or []
    collected: dict[str, dict[str, Any]] = {}

    def fetch_campaign(concept: dict[str, Any]) -> tuple[str, list[dict[str, Any]]]:
        concept_id = concept.get("conceptos_id")
        campaign = clean(concept.get("nombre"))
        result: list[dict[str, Any]] = []
        page, page_size = 1, 100
        while True:
            payload = client.get(
                "getBeneficiosSql",
                idBenefConcepto=concept_id,
                page=page,
                page_size=page_size,
            )
            items = payload.get("data") or []
            result.extend(items)
            total = int(payload.get("total") or len(items))
            if not items or page * page_size >= total:
                break
            page += 1
        return campaign, result

    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(fetch_campaign, c): c for c in concepts}
        for number, future in enumerate(as_completed(futures), 1):
            campaign, items = future.result()
            for item in items:
                key = str(item.get("id") or promo_url(item))
                if key not in collected:
                    collected[key] = dict(item)
                    collected[key]["_campanias"] = []
                if campaign and campaign not in collected[key]["_campanias"]:
                    collected[key]["_campanias"].append(campaign)
            print(f"  Campaña {number}/{len(concepts)}: {campaign}", flush=True)
    return list(collected.values())


def fetch_complete(
    client: ClubHolaClient, workers: int, with_details: bool = True
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    page, page_size = 1, 100
    while True:
        payload = client.get("getBuscarBeneficios", page=page, page_size=page_size)
        current = payload.get("data") or []
        items.extend(current)
        total = int(payload.get("total") or len(items))
        print(f"  Catálogo: {min(len(items), total)}/{total}", flush=True)
        if not current or len(items) >= total:
            break
        page += 1

    unique = {str(x.get("id")): x for x in items if x.get("id") is not None}
    if not with_details:
        output = list(unique.values())
        for item in output:
            item["_detalle_ok"] = False
            item["_detalle_error"] = "Detalle omitido por --sin-detalle"
        return output

    def detail(item: dict[str, Any]) -> dict[str, Any]:
        try:
            data = client.get("detalleBeneficio", idBeneficio=item["id"]).get("data") or {}
            merged = {**item, **data}
            merged["_detalle_ok"] = True
            return merged
        except Exception as exc:  # conserva la fila aunque falle una ficha
            item["_detalle_ok"] = False
            item["_detalle_error"] = str(exc)
            return item

    output: list[dict[str, Any]] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = [pool.submit(detail, x) for x in unique.values()]
        for count, future in enumerate(as_completed(futures), 1):
            output.append(future.result())
            if count % 25 == 0 or count == len(futures):
                print(f"  Detalles: {count}/{len(futures)}", flush=True)
    return output


def request_dni() -> str:
    """Lee el DNI exclusivamente desde CLUBHOLA_DNI; nunca se guarda en el código."""
    dni = os.getenv("CLUBHOLA_DNI", DEFAULT_DNI).strip()
    if not re.fullmatch(r"\d{8}", dni):
        raise ValueError("Defina CLUBHOLA_DNI como secreto con exactamente 8 dígitos.")
    print("DNI configurado automáticamente.", flush=True)
    return dni


def card_to_item(text: str, url: str) -> dict[str, Any]:
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    categories = {
        "gastronomía", "entretenimiento", "hogar", "wellness", "mascotas",
        "turismo", "moda", "belleza", "educación",
    }
    category = lines[-1] if lines and lines[-1].lower() in categories else ""
    name_index = -2 if category and len(lines) >= 2 else -1
    name = lines[name_index] if lines else ""
    benefit_lines = lines[:name_index] if name_index else []
    promo_id = ""
    match = re.search(r"-(\d+)(?:[/?#]|$)", url)
    if match:
        promo_id = match.group(1)
    return {
        "id": promo_id,
        "titulo": name,
        "nombre": category,
        "chapita": " ".join(benefit_lines),
        "url_beneficio": url,
        "_detalle_ok": None,
    }


def section_between(lines: list[str], start: str, stops: tuple[str, ...]) -> str:
    try:
        position = next(i for i, value in enumerate(lines) if value == start) + 1
    except StopIteration:
        return ""
    result: list[str] = []
    for value in lines[position:]:
        if value in stops:
            break
        result.append(value)
    return "\n".join(result).strip()


def parse_browser_detail(body_text: str) -> dict[str, Any]:
    """Transforma los encabezados visibles de la ficha en campos del Excel."""
    lines = [clean(x) for x in body_text.splitlines() if clean(x)]
    legal_stop = (
        "Beneficios similares", "Términos y condiciones", "Políticas de privacidad",
        "Preguntas frecuentes", "TODOS LOS DERECHOS RESERVADOS © 2026 El Comercio",
    )
    return {
        "descripcion": section_between(
            lines, "Conoce como disfrutar tu beneficio", ("Ubicación",)
        ),
        "direccion": section_between(lines, "Ubicación", ("Vigencia",)),
        "cuando": section_between(lines, "Vigencia", ("Cómo usar la promoción",)),
        "como": section_between(
            lines,
            "Cómo usar la promoción",
            ("Términos y condiciones legales", "Beneficios similares"),
        ),
        "terminos_condiciones_web": section_between(
            lines, "Términos y condiciones legales", legal_stop
        ),
    }


def find_access_token(value: Any) -> str:
    """Busca access_token dentro de una respuesta JSON sin mostrarlo ni guardarlo."""
    if isinstance(value, dict):
        token = value.get("access_token")
        if isinstance(token, str) and token.strip():
            return token.strip()
        for child in value.values():
            found = find_access_token(child)
            if found:
                return found
    elif isinstance(value, list):
        for child in value:
            found = find_access_token(child)
            if found:
                return found
    return ""


def capture_browser_token(driver: Any, seconds: float = 8.0) -> str:
    """Recupera la sesión efímera desde el tráfico de Chrome.

    El valor se mantiene únicamente en memoria: nunca se imprime ni se escribe
    en el Excel. Primero revisa peticiones autenticadas y, como respaldo, la
    respuesta del inicio de sesión.
    """
    deadline = time.time() + seconds
    response_request_ids: list[str] = []
    request_urls: dict[str, str] = {}

    def token_from_headers(headers: Any) -> str:
        if not isinstance(headers, dict):
            return ""
        for name, value in headers.items():
            if name.lower() == "authorization" and str(value).startswith("Bearer "):
                return str(value)[7:].strip()
        return ""

    while time.time() < deadline:
        try:
            entries = driver.get_log("performance")
        except Exception:
            return ""
        for entry in entries:
            try:
                event = json.loads(entry["message"])["message"]
                method = event.get("method")
                params = event.get("params") or {}
                if method == "Network.requestWillBeSent":
                    request = params.get("request") or {}
                    request_url = clean(request.get("url"))
                    request_id = clean(params.get("requestId"))
                    if request_id:
                        request_urls[request_id] = request_url
                    headers = request.get("headers") or {}
                    if "/apimovil/" in request_url:
                        token = token_from_headers(headers)
                        if token:
                            return token
                elif method == "Network.requestWillBeSentExtraInfo":
                    request_id = clean(params.get("requestId"))
                    if "/apimovil/" in request_urls.get(request_id, ""):
                        token = token_from_headers(params.get("headers"))
                        if token:
                            return token
                elif method == "Network.responseReceived":
                    response = params.get("response") or {}
                    if "/loginByDocument" in clean(response.get("url")):
                        response_request_ids.append(clean(params.get("requestId")))
            except (KeyError, TypeError, ValueError):
                continue

        for request_id in response_request_ids:
            if not request_id:
                continue
            try:
                result = driver.execute_cdp_cmd(
                    "Network.getResponseBody", {"requestId": request_id}
                )
                raw = result.get("body", "")
                if result.get("base64Encoded"):
                    raw = base64.b64decode(raw).decode("utf-8", errors="replace")
                token = find_access_token(json.loads(raw))
                if token:
                    return token
            except Exception:
                continue
        response_request_ids.clear()
        time.sleep(0.25)
    return ""


def fetch_with_dni(dni: str, args: argparse.Namespace) -> list[dict[str, Any]]:
    """Inicia sesión y recorre el catálogo que ClubHOLA muestra al usuario."""
    try:
        from selenium import webdriver
        from selenium.common.exceptions import TimeoutException
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support import expected_conditions as EC
        from selenium.webdriver.support.ui import Select, WebDriverWait
    except ImportError as exc:
        raise RuntimeError(
            "Falta Selenium. Instálelo con: pip install selenium"
        ) from exc

    options = webdriver.ChromeOptions()
    if args.mostrar_navegador:
        options.add_argument("--start-maximized")
    else:
        # Chrome funciona en segundo plano; conserva el mismo tamaño del
        # formulario de escritorio para que se rendericen los controles válidos.
        options.add_argument("--headless=new")
        options.add_argument("--window-size=1440,1100")
        options.add_argument("--disable-gpu")
    options.add_argument("--disable-notifications")
    # No hace falta esperar imágenes y recursos secundarios para empezar a
    # completar el formulario. Esto reduce bastante la demora inicial.
    options.page_load_strategy = "eager"
    options.set_capability("goog:loggingPrefs", {"performance": "ALL"})
    driver = webdriver.Chrome(options=options)  # Selenium Manager obtiene el driver.
    driver.set_page_load_timeout(45)
    try:
        driver.execute_cdp_cmd("Network.enable", {})
    except Exception:
        pass  # El recorrido visual seguirá disponible como respaldo.
    wait = WebDriverWait(driver, 45)

    try:
        login_url = f"{HOME}/sign-in?continue=%2Fcatalogo-virtual"
        driver.get(login_url)
        WebDriverWait(driver, 30).until(
            lambda d: d.execute_script("return document.readyState") in ("interactive", "complete")
        )

        def session_open() -> bool:
            # La señal principal es "Cerrar sesión". Algunas versiones de la
            # aplicación conservan temporalmente /sign-in en la URL aunque la
            # cuenta ya haya sido reconocida.
            if any(
                element.is_displayed()
                for element in driver.find_elements(
                    By.XPATH, "//*[normalize-space()='Cerrar sesión']"
                )
            ):
                return True
            current = driver.current_url.lower()
            return "catalogo-virtual" in current and "sign-in" not in current

        session_token = ""
        if not session_open():
            automatic_login = False
            try:
                # Se escoge el select que realmente contiene la opción DNI; así se
                # evita confundirlo con filtros del menú si la página cambia.
                document_type = WebDriverWait(driver, 25).until(
                    lambda d: next(
                        (
                            select
                            for select in d.find_elements(By.TAG_NAME, "select")
                            if select.is_displayed()
                            and select.is_enabled()
                            and any(
                                option.get_attribute("value") == "DNI"
                                for option in select.find_elements(By.TAG_NAME, "option")
                            )
                        ),
                        False,
                    )
                )
                try:
                    Select(document_type).select_by_value("DNI")
                except Exception:
                    # Algunos componentes de React muestran un select nativo
                    # transparente. En ese caso se cambia su valor directamente
                    # y se emite un solo evento change para actualizar React.
                    driver.execute_script(
                        """
                        const select = arguments[0];
                        select.value = 'DNI';
                        select.dispatchEvent(new Event('change', {bubbles: true}));
                        """,
                        document_type,
                    )

                def dni_field(d):
                    candidates = d.find_elements(
                        By.CSS_SELECTOR,
                        "input:not([type='hidden']):not([type='submit']), textarea",
                    )
                    visible = [
                        field for field in candidates
                        if field.is_displayed() and field.is_enabled()
                    ]
                    if not visible:
                        return False

                    def score(field):
                        attributes = normalized(
                            " ".join(
                                filter(
                                    None,
                                    (
                                        field.get_attribute("placeholder"),
                                        field.get_attribute("aria-label"),
                                        field.get_attribute("name"),
                                        field.get_attribute("id"),
                                        field.get_attribute("inputmode"),
                                    ),
                                )
                            )
                        )
                        points = 8 if any(
                            word in attributes
                            for word in ("documento", "numero", "dni")
                        ) else 0
                        points += 4 if field.get_attribute("maxlength") == "8" else 0
                        points += 2 if "numeric" in attributes else 0
                        return points

                    return max(visible, key=score)

                dni_input = WebDriverWait(driver, 20).until(dni_field)
                typed_normally = False
                try:
                    dni_input.click()
                    dni_input.clear()
                    dni_input.send_keys(dni)
                    WebDriverWait(driver, 5).until(
                        lambda _d: dni_input.get_attribute("value") == dni
                    )
                    typed_normally = True
                except Exception:
                    pass

                if not typed_normally:
                    # Respaldo para componentes que impiden clear/send_keys.
                    driver.execute_script(
                        """
                        const el = arguments[0];
                        const setter = Object.getOwnPropertyDescriptor(
                            HTMLInputElement.prototype, 'value'
                        ).set;
                        setter.call(el, arguments[1]);
                        el.dispatchEvent(new Event('input', {bubbles: true}));
                        el.dispatchEvent(new Event('change', {bubbles: true}));
                        """,
                        dni_input,
                        dni,
                    )

                # ClubHOLA renderiza dos controles llamados "Ingresar": uno está
                # oculto para otro tamaño de pantalla. Se escoge expresamente el
                # DIV visible y habilitado que recibe el clic real.
                login_xpath = (
                    "//*[normalize-space()='Ingresar']/ancestor-or-self::*"
                    "[@tabindex='0' or self::button][1]"
                )

                def visible_login_control(d):
                    return next(
                        (
                            element
                            for element in d.find_elements(By.XPATH, login_xpath)
                            if element.is_displayed() and element.is_enabled()
                        ),
                        False,
                    )

                login_control = WebDriverWait(driver, 15).until(
                    visible_login_control
                )
                driver.execute_script(
                    "arguments[0].scrollIntoView({block: 'center'});", login_control
                )
                time.sleep(0.4)
                try:
                    login_control.click()
                except Exception:
                    driver.execute_script(
                        """
                        const el = arguments[0];
                        el.dispatchEvent(new PointerEvent('pointerdown', {bubbles: true}));
                        el.dispatchEvent(new MouseEvent('mousedown', {bubbles: true}));
                        el.dispatchEvent(new PointerEvent('pointerup', {bubbles: true}));
                        el.dispatchEvent(new MouseEvent('mouseup', {bubbles: true}));
                        el.click();
                        """,
                        login_control,
                    )
                print("Botón Ingresar presionado automáticamente.", flush=True)
                automatic_login = True
            except Exception as exc:
                print(
                    "ClubHOLA cambió o demoró en completar el ingreso automático "
                    f"({type(exc).__name__}).",
                    flush=True,
                )

            if automatic_login:
                # El token aparece en la respuesta del login antes de que React
                # termine de cambiar la pantalla. Detectarlo aquí evita esperar
                # innecesariamente hasta 90 segundos.
                session_token = capture_browser_token(driver, seconds=15)
                if not session_token:
                    try:
                        WebDriverWait(driver, 30).until(lambda _d: session_open())
                    except TimeoutException:
                        pass

            if not session_open() and not session_token:
                # Comprobación automática adicional: una sesión recién creada
                # puede estar lista antes de que /sign-in cambie visualmente.
                driver.get(f"{HOME}/catalogo-virtual/")
                try:
                    WebDriverWait(driver, 30).until(lambda _d: session_open())
                except TimeoutException:
                    pass
                session_token = capture_browser_token(driver, seconds=3)

        if not session_open() and not session_token:
            diagnostic = Path.cwd() / "clubhola_diagnostico.png"
            try:
                driver.save_screenshot(str(diagnostic))
                diagnostic_message = f" Captura guardada en: {diagnostic}"
            except Exception:
                diagnostic_message = ""
            raise RuntimeError(
                "ClubHOLA no confirmó el ingreso automático. Si el portal está "
                "mostrando una verificación, vuelva a ejecutar con "
                "--mostrar-navegador." + diagnostic_message
            )
        print("Sesión ClubHOLA reconocida correctamente.", flush=True)

        # Primer intento: normalmente el token está en la respuesta del login o
        # en las peticiones que carga la portada después de ingresar.
        if not session_token:
            session_token = capture_browser_token(driver, seconds=2.5)

        if session_token:
            print(
                "Sesión temporal detectada. Descargando catálogo mediante API...",
                flush=True,
            )
            try:
                return fetch_complete(
                    ClubHolaClient(session_token),
                    args.workers,
                    with_details=not args.sin_detalle,
                )
            except Exception as exc:
                print(
                    "La API autenticada no respondió correctamente; se continuará "
                    f"con el navegador como respaldo ({type(exc).__name__}).",
                    flush=True,
                )
                session_token = ""

        first_page = f"{HOME}/catalogo-virtual/page/1?orden=desc"
        driver.get(first_page)
        wait.until(EC.presence_of_element_located((By.XPATH, "//*[contains(.,'Ordenar por:')]")))
        WebDriverWait(driver, 45).until(
            lambda d: len(d.find_elements(By.CSS_SELECTOR, "a[href]")) > 10
        )
        time.sleep(1.5)

        # Segundo intento: el catálogo siempre realiza peticiones autenticadas.
        if not session_token:
            session_token = capture_browser_token(driver, seconds=3)

        if session_token:
            print(
                "Sesión temporal detectada. Descargando catálogo mediante API...",
                flush=True,
            )
            try:
                return fetch_complete(
                    ClubHolaClient(session_token),
                    args.workers,
                    with_details=not args.sin_detalle,
                )
            except Exception as exc:
                print(
                    "La API autenticada no respondió correctamente; se continuará "
                    f"con el navegador como respaldo ({type(exc).__name__}).",
                    flush=True,
                )
        else:
            print(
                "No fue posible recuperar la sesión temporal; se continuará con "
                "el navegador como respaldo.",
                flush=True,
            )

        page_numbers = []
        for element in driver.find_elements(By.CSS_SELECTOR, "[tabindex='0']"):
            value = clean(element.text)
            if value.isdigit():
                page_numbers.append(int(value))
        total_pages = max(page_numbers, default=1)
        print(f"Páginas del catálogo detectadas: {total_pages}", flush=True)

        promo_pattern = re.compile(
            rf"^{re.escape(HOME)}/[a-z0-9áéíóúñü-]+-\d+(?:[/?#].*)?$", re.I
        )
        collected: dict[str, dict[str, Any]] = {}
        for page in range(1, total_pages + 1):
            url = f"{HOME}/catalogo-virtual/page/{page}?orden=desc"
            if driver.current_url != url:
                driver.get(url)
            WebDriverWait(driver, 45).until(
                lambda d: "Ordenar por:" in d.find_element(By.TAG_NAME, "body").text
            )
            time.sleep(args.espera)
            for anchor in driver.find_elements(By.CSS_SELECTOR, "a[href]"):
                href = clean(anchor.get_attribute("href"))
                if not promo_pattern.match(href):
                    continue
                item = card_to_item(anchor.text, href)
                key = str(item.get("id") or href)
                existing = collected.get(key)
                if existing is None or len(clean(item.get("titulo"))) > len(clean(existing.get("titulo"))):
                    collected[key] = item
            print(
                f"  Página {page}/{total_pages}: {len(collected)} promociones únicas",
                flush=True,
            )

        items = list(collected.values())
        # Estos filtros reducen el número de fichas que hay que abrir.
        if args.categoria:
            needle = normalized(args.categoria)
            items = [x for x in items if needle in normalized(x.get("nombre"))]
        if args.buscar:
            needle = normalized(args.buscar)
            items = [x for x in items if needle in normalized(f"{x.get('titulo')} {x.get('chapita')}")]
        if args.limit:
            items = items[: args.limit]

        if args.sin_detalle:
            for item in items:
                item["_detalle_ok"] = False
                item["_detalle_error"] = "Detalle omitido por --sin-detalle"
            return items

        for number, item in enumerate(items, 1):
            try:
                driver.get(item["url_beneficio"])
                WebDriverWait(driver, 45).until(
                    lambda d: "Conoce como disfrutar tu beneficio"
                    in d.find_element(By.TAG_NAME, "body").text
                )
                detail = parse_browser_detail(driver.find_element(By.TAG_NAME, "body").text)
                item.update(detail)
                item["_detalle_ok"] = True
            except Exception as exc:
                item["_detalle_ok"] = False
                item["_detalle_error"] = f"No se pudo leer la ficha: {exc}"
            if number % 20 == 0 or number == len(items):
                print(f"  Fichas detalladas: {number}/{len(items)}", flush=True)
        return items
    finally:
        if not args.mantener_navegador:
            driver.quit()


def parse_date(value: Any) -> datetime | None:
    text = clean(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def validity(item: dict[str, Any]) -> str:
    now = datetime.now()
    start = parse_date(item.get("fecha_inicio_vigencia"))
    end = parse_date(item.get("fecha_fin_vigencia"))
    if not truthy(item.get("activo", 1)):
        return "Inactiva"
    if truthy(item.get("sin_stock")):
        return "Sin stock"
    if start and now < start:
        return "Próxima"
    if end and now > end:
        return "Vencida"
    return "Vigente"


def to_row(item: dict[str, Any], coverage: str, extracted_at: str) -> dict[str, Any]:
    status = "OK"
    if item.get("_detalle_ok") is False:
        status = "Ficha parcial: " + clean(item.get("_detalle_error"))
    detail = clean(item.get("descripcion") or item.get("descripcion_corta"))
    branches = join_dicts(item.get("sucursales"))
    return {
        "Banco": "BanBif",
        "Tipo": "Promoción / beneficio ClubHOLA",
        "Nombre": clean(item.get("titulo")),
        "Comercio": clean(item.get("est_nombre")),
        "Detalle de la promoción": detail,
        "Cómo obtener el descuento": clean(item.get("como")),
        "Restricciones / vigencia": clean(item.get("cuando")),
        "Términos y condiciones": clean(item.get("terminos_condiciones_web")),
        "Ubicación": clean(item.get("direccion")),
        "Sucursales": branches,
        "Link": promo_url(item),
        "Enlace online": clean(item.get("url_online")),
        "Web del comercio": clean(item.get("web")),
        "Categoría": clean(item.get("nombre") or item.get("tipo_beneficio_nombre")),
        "Subcategoría": clean(item.get("cat_nombre")),
        "Campañas": clean(item.get("_campanias") or item.get("nombre_campania")),
        "ID promoción": item.get("id", ""),
        "ID comercio": item.get("establecimiento_id", ""),
        "Beneficio": clean(item.get("chapita") or item.get("valor")),
        "Prefijo beneficio": clean(item.get("chapita_texto_superior")),
        "Sufijo beneficio": clean(item.get("chapita_texto_inferior")),
        "Fecha de inicio": clean(item.get("fecha_inicio_vigencia")),
        "Fecha de fin": clean(item.get("fecha_fin_vigencia")),
        "Estado de vigencia": validity(item),
        "Solo por hoy": "Sí" if truthy(item.get("solo_por_hoy")) else "No",
        "Niveles ClubHOLA": clean(item.get("tipo_suscriptor_nombre") or item.get("tipo_suscriptor_codigo")),
        "Teléfono": clean(item.get("telefono_info")),
        "Email": clean(item.get("email_info")),
        "Stock total": item.get("stock", ""),
        "Stock disponible": item.get("stock_actual", ""),
        "Sin stock": "Sí" if truthy(item.get("sin_stock")) else "No",
        "Requiere cupón": "Sí" if truthy(item.get("generar_cupon")) else "No",
        "Compra con login": "Sí" if truthy(item.get("compra_con_login")) else "No",
        "Máximo por suscriptor": item.get("maximo_por_subscriptor", item.get("maximo_por_subscriptoras", "")),
        "Valoración promedio": item.get("valoracion_prom", ""),
        "Cantidad de valoraciones": item.get("valoracion_cant", ""),
        "Imagen": clean(item.get("path_logo")),
        "Fecha de actualización": clean(item.get("fecha_actualizacion")),
        "Cobertura": coverage,
        "Estado de extracción": status,
        "Fecha extracción UTC": extracted_at,
    }


def apply_filters(items: list[dict[str, Any]], args: argparse.Namespace) -> list[dict[str, Any]]:
    output = items
    if args.categoria:
        needle = normalized(args.categoria)
        output = [
            x for x in output
            if needle in normalized(" ".join(clean(x.get(k)) for k in (
                "nombre", "tipo_beneficio_nombre", "tbenef_slug", "cat_nombre", "cat_slug"
            )))
        ]
    if args.buscar:
        needle = normalized(args.buscar)
        output = [x for x in output if needle in normalized(" ".join(clean(v) for v in x.values()))]
    if args.solo_vigentes:
        output = [x for x in output if validity(x) == "Vigente"]
    if args.solo_disponibles:
        output = [x for x in output if not truthy(x.get("sin_stock"))]
    if args.limit:
        output = output[: args.limit]
    return output


def save_excel(rows: list[dict[str, Any]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows, columns=COLUMNS)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Promociones")
        ws = writer.book["Promociones"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        fill = PatternFill("solid", fgColor="123B70")
        for cell in ws[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths = {
            "Nombre": 38, "Detalle de la promoción": 55, "Cómo obtener el descuento": 48,
            "Restricciones / vigencia": 42, "Términos y condiciones": 65,
            "Ubicación": 42, "Sucursales": 55, "Link": 48, "Imagen": 48,
            "Campañas": 35, "Estado de extracción": 34,
        }
        for index, column in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(index)].width = widths.get(column, 20)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[1].height = 32


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Scraper BanBif / ClubHOLA a Excel")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Ruta del archivo .xlsx")
    parser.add_argument(
        "--modo", choices=("dni", "public", "token"), default="dni",
        help="dni usa Chrome oculto; public no inicia sesión; token usa CLUBHOLA_TOKEN",
    )
    parser.add_argument("--categoria", help="Categoría o subcategoría (p. ej. gastronomia)")
    parser.add_argument("--buscar", help="Texto que debe aparecer en la promoción")
    parser.add_argument("--solo-vigentes", action="store_true")
    parser.add_argument("--solo-disponibles", action="store_true")
    parser.add_argument("--workers", type=int, default=4, help="Hilos para detalles autenticados")
    parser.add_argument("--limit", type=int, default=0, help="Límite de filas (0 = todas)")
    parser.add_argument(
        "--sin-detalle", action="store_true",
        help="En modo dni exporta las tarjetas sin abrir cada ficha (mucho más rápido)",
    )
    parser.add_argument(
        "--espera", type=float, default=1.5,
        help="Segundos adicionales por página del catálogo en modo dni",
    )
    parser.add_argument(
        "--mantener-navegador", action="store_true",
        help="No cierra Chrome al terminar",
    )
    parser.add_argument(
        "--mostrar-navegador", action="store_true",
        help="Muestra Chrome; de forma predeterminada se ejecuta oculto",
    )
    return parser.parse_args()


def main() -> int:
    args = arguments()
    token = os.getenv("CLUBHOLA_TOKEN", "").strip()
    if args.modo == "token" and not token:
        print(
            "ERROR: el modo token requiere CLUBHOLA_TOKEN. "
            "ClubHOLA protege el catálogo y el detalle detrás del inicio de sesión.",
            file=sys.stderr,
        )
        return 2

    client = ClubHolaClient(token)
    print(f"Consultando ClubHOLA en modo {args.modo}...", flush=True)
    if args.modo == "dni":
        items = fetch_with_dni(request_dni(), args)
        coverage = "Catálogo completo autenticado mediante DNI"
    elif args.modo == "token":
        items = fetch_complete(client, args.workers)
        coverage = "Catálogo completo autenticado + fichas de detalle"
    else:
        items = fetch_public(client, args.workers)
        coverage = "Campañas públicas visibles sin iniciar sesión"

    print(f"Promociones únicas obtenidas: {len(items)}", flush=True)
    items = apply_filters(items, args)
    extracted_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    rows = [to_row(x, coverage, extracted_at) for x in items]
    rows.sort(key=lambda x: (x["Categoría"], x["Nombre"]))
    save_excel(rows, Path(args.output).expanduser().resolve())
    print(f"Filas exportadas: {len(rows)}")
    print(f"Excel creado: {Path(args.output).expanduser().resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
