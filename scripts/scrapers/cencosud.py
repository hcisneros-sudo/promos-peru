#!/usr/bin/env python3
"""Extrae promociones de Tarjeta Cencosud Perú sin Selenium.

Estrategia:
1. Descubre el tipo de contenido en WordPress REST y descarga sus registros.
2. Si REST no está expuesto, obtiene las URLs desde el sitemap de WordPress.
3. Como último respaldo, descubre las fichas desde el catálogo HTML.

Instalación:
    pip install requests beautifulsoup4 pandas openpyxl

Uso:
    python Cencosud_Promociones_API.py
    python Cencosud_Promociones_API.py --solo-vigentes
    python Cencosud_Promociones_API.py --limit 20
    python Cencosud_Promociones_API.py --fuente api
"""

from __future__ import annotations

import argparse
import hashlib
import html
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup, Tag
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


BASE_URL = "https://www.tarjetacencosud.pe"
CATALOG_URL = f"{BASE_URL}/promociones/"
DEFAULT_OUTPUT = "Cencosud_Promociones.xlsx"
TIMEOUT = 45

HEADERS = {
    "Accept": "text/html,application/xhtml+xml,application/json;q=0.9,*/*;q=0.8",
    "Accept-Language": "es-PE,es;q=0.9,en;q=0.7",
    "Cache-Control": "no-cache",
    "Referer": f"{BASE_URL}/",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/152.0.0.0 Safari/537.36"
    ),
}

COLUMNS = [
    "Banco", "Tipo", "Nombre", "Detalle de la promoción", "Restricciones",
    "Términos y condiciones", "Ubicación", "Link", "Categorías",
    "Establecimiento", "ID promoción", "Descuento (%)", "Texto del beneficio",
    "Precio promocional", "Precio regular", "Cuotas sin intereses",
    "Fecha de inicio", "Fecha de fin", "Estado de vigencia", "Días válidos",
    "Tarjetas válidas", "Regiones / zonas", "Comercio", "Código de cupón",
    "Enlaces relacionados", "Imagen", "Fuente técnica", "Estado de extracción",
    "Fecha de extracción (UTC)",
]

MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
    "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9,
    "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}

KNOWN_LOCATIONS = [
    "Nivel Nacional", "Lima", "Callao", "Arequipa", "Puno", "Tumbes",
    "Ayacucho", "Áncash", "Ancash", "Huánuco", "San Martín", "Moquegua",
    "Cajamarca", "Trujillo", "Loreto", "Cusco", "Junín", "La Libertad",
    "Lambayeque", "Lambayaque", "Ica", "Piura", "Tacna", "Ucayali",
    "Uyacali", "Ancasch",
]


def clean_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (list, tuple, set)):
        return " | ".join(x for x in (clean_text(v) for v in value) if x)
    text = html.unescape(str(value)).replace("\xa0", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def slugify(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean_text(value).lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def unique(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    for value in values:
        text = clean_text(value)
        if text and text not in result:
            result.append(text)
    return result


def unique_join(values: Iterable[Any], separator: str = " | ") -> str:
    return separator.join(unique(values))


def html_text(fragment: Any) -> str:
    soup = BeautifulSoup(clean_text(fragment), "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    return clean_text(soup.get_text("\n", strip=True))


def make_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=3,
        connect=3,
        read=3,
        backoff_factor=0.7,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset(("GET", "HEAD")),
    )
    session.mount("https://", HTTPAdapter(max_retries=retry, pool_connections=20, pool_maxsize=20))
    session.headers.update(HEADERS)
    return session


def checked_get(
    session: requests.Session,
    url: str,
    *,
    params: dict[str, Any] | None = None,
    timeout: int = TIMEOUT,
) -> requests.Response:
    response = session.get(url, params=params, timeout=timeout)
    body_start = response.text[:3000].lower()
    if "sorry, you have been blocked" in body_start or "attention required" in body_start:
        raise RuntimeError(
            "Cloudflare bloqueó la conexión HTTP. Pruebe desde su conexión habitual "
            "o ejecute nuevamente más tarde."
        )
    response.raise_for_status()
    return response


def find_rest_base(session: requests.Session) -> str:
    """Devuelve el rest_base del post type de promociones, si está expuesto."""
    try:
        response = checked_get(session, f"{BASE_URL}/wp-json/wp/v2/types")
        types = response.json()
        for key, item in types.items():
            signature = slugify(
                f"{key} {item.get('slug', '')} {item.get('rest_base', '')} "
                f"{item.get('name', '')}"
            )
            if "promot" in signature or "promoc" in signature:
                return clean_text(item.get("rest_base") or item.get("slug") or key)
    except (requests.RequestException, ValueError, RuntimeError):
        pass

    # Algunos WordPress ocultan /types, pero mantienen accesible el endpoint
    # del custom post type. Se prueban solo los nombres convencionales.
    for candidate in ("promotion", "promotions", "promociones"):
        try:
            probe = checked_get(
                session,
                f"{BASE_URL}/wp-json/wp/v2/{candidate}",
                params={"per_page": 1, "page": 1},
            ).json()
            if isinstance(probe, list):
                return candidate
        except (requests.RequestException, ValueError, RuntimeError):
            continue
    return ""


def embedded_terms(post: dict[str, Any]) -> list[str]:
    groups = post.get("_embedded", {}).get("wp:term", [])
    return unique(
        term.get("name")
        for group in groups
        if isinstance(group, list)
        for term in group
        if isinstance(term, dict)
    )


def featured_image(post: dict[str, Any]) -> str:
    media = post.get("_embedded", {}).get("wp:featuredmedia", [])
    if media and isinstance(media[0], dict):
        return clean_text(media[0].get("source_url"))
    return ""


def fetch_rest_posts(
    session: requests.Session,
    rest_base: str,
    limit: int = 0,
) -> list[dict[str, Any]]:
    endpoint = f"{BASE_URL}/wp-json/wp/v2/{rest_base.strip('/')}"
    output: list[dict[str, Any]] = []
    page = 1
    while True:
        response = checked_get(
            session,
            endpoint,
            params={"per_page": 100, "page": page, "status": "publish", "_embed": 1},
        )
        batch = response.json()
        if not isinstance(batch, list):
            raise RuntimeError("WordPress REST no devolvió una lista de promociones.")
        output.extend(batch)
        total_pages = int(response.headers.get("X-WP-TotalPages", page))
        print(f"  API WordPress: página {page}/{total_pages} ({len(output)} registros)", flush=True)
        if (limit and len(output) >= limit) or page >= total_pages or not batch:
            break
        page += 1
    return output[:limit] if limit else output


def xml_locations(xml: str) -> list[str]:
    soup = BeautifulSoup(xml, "xml")
    return unique(node.get_text(strip=True) for node in soup.find_all("loc"))


def discover_from_sitemap(session: requests.Session) -> list[str]:
    indexes: list[str] = []
    for candidate in (f"{BASE_URL}/wp-sitemap.xml", f"{BASE_URL}/sitemap_index.xml"):
        try:
            indexes = xml_locations(checked_get(session, candidate).text)
            if indexes:
                break
        except (requests.RequestException, RuntimeError):
            continue

    urls: list[str] = []
    for location in indexes:
        path = urlparse(location).path.lower()
        if "/promotion/" in path:
            urls.append(location)
        elif "promotion" in path or "promoc" in path:
            try:
                urls.extend(
                    url for url in xml_locations(checked_get(session, location).text)
                    if "/promotion/" in urlparse(url).path.lower()
                )
            except (requests.RequestException, RuntimeError):
                continue
    return unique(urls)


def discover_from_catalog(session: requests.Session) -> list[str]:
    soup = BeautifulSoup(checked_get(session, CATALOG_URL).text, "html.parser")
    return unique(
        urljoin(BASE_URL, anchor.get("href"))
        for anchor in soup.select("a[href]")
        if "/promotion/" in urlparse(urljoin(BASE_URL, anchor.get("href"))).path.lower()
    )


def section_from_heading(soup: BeautifulSoup, name: str) -> str:
    target = slugify(name)
    headings = soup.find_all(re.compile(r"^h[1-6]$"))
    heading = next((item for item in headings if slugify(item.get_text(" ", strip=True)) == target), None)
    if not heading:
        return ""
    level = int(heading.name[1])
    parts: list[str] = []
    for node in heading.next_elements:
        if node is heading:
            continue
        if isinstance(node, Tag) and re.fullmatch(r"h[1-6]", node.name or ""):
            if int(node.name[1]) <= level:
                break
        if isinstance(node, Tag) and node.name in ("p", "li"):
            text = clean_text(node.get_text(" ", strip=True))
            if text:
                parts.append(text)
    return unique_join(parts, "\n")


def nearest_heading_before(reference: Tag | None, level: str) -> str:
    if not reference:
        return ""
    previous = reference.find_all_previous(level)
    return clean_text(previous[0].get_text(" ", strip=True)) if previous else ""


def visible_dates(text: str) -> list[date]:
    source = clean_text(text)
    normalized = slugify(source).replace("-", " ")
    found: list[tuple[int, int, int, int | None]] = []
    for match in re.finditer(r"(?<!\d)(\d{1,2})[/-](\d{1,2})[/-](20\d{2})(?!\d)", source):
        found.append((match.start(), int(match.group(1)), int(match.group(2)), int(match.group(3))))

    months = "|".join(MONTHS)
    textual: list[tuple[int, int, int, int | None]] = []
    for match in re.finditer(
        rf"(?<!\d)(\d{{1,2}})\s+(?:de\s+)?({months})(?:\s+(?:de|del)\s+(20\d{{2}}))?",
        normalized,
    ):
        textual.append((
            match.start(), int(match.group(1)), MONTHS[match.group(2)],
            int(match.group(3)) if match.group(3) else None,
        ))
    explicit = [(position, year) for position, _, _, year in found + textual if year]
    for position, day, month, year in textual:
        if year is None and explicit:
            year = min(explicit, key=lambda item: abs(item[0] - position))[1]
        found.append((position, day, month, year))

    result: list[date] = []
    for _, day, month, year in sorted(found):
        if year:
            try:
                parsed = date(year, month, day)
                if parsed not in result:
                    result.append(parsed)
            except ValueError:
                pass
    return result


def validity(start: date | None, end: date | None) -> str:
    today = date.today()
    if start and today < start:
        return "Programada"
    if end and today > end:
        return "Finalizada (aún publicada)"
    if start or end:
        return "Vigente"
    return "Publicada (fecha no identificada)"


def percentage(text: str) -> int | float | str:
    match = re.search(r"(?<!\d)(\d{1,3}(?:[.,]\d+)?)\s*%", clean_text(text))
    if not match:
        return ""
    number = float(match.group(1).replace(",", "."))
    return int(number) if number.is_integer() else number


def money_values(text: str) -> tuple[str, str]:
    source = clean_text(text)
    values = re.findall(r"(?:S/\s*|US\$\s*|\$\s*)(\d+(?:[.,]\d{1,2})?)", source, re.I)
    regular_match = re.search(
        r"precio\s+regular\s*:?[ \t]*(?:S/\s*|US\$\s*|\$\s*)?(\d+(?:[.,]\d{1,2})?)",
        source, re.I,
    )
    regular = regular_match.group(1).replace(",", ".") if regular_match else ""
    promo = next((value.replace(",", ".") for value in values if value.replace(",", ".") != regular), "")
    return promo, regular


def coupon_code(text: str) -> str:
    return unique_join(
        match.upper()
        for match in re.findall(
            r"(?:c[oó]digo|cup[oó]n|nombre\s+del\s+cup[oó]n)\s*:?[ \t]*([A-Z0-9-]{4,30})",
            clean_text(text), re.I,
        )
    )


def installments(text: str) -> str:
    values: list[str] = []
    for match in re.findall(r"((?:\d+[ ,y]*)+)\s+cuotas?\s+sin\s+intereses", clean_text(text), re.I):
        values.extend(re.findall(r"\d+", match))
    return unique_join(values, ", ")


def valid_days(text: str) -> str:
    canonical = {
        "lunes": "Lunes", "martes": "Martes", "miércoles": "Miércoles",
        "miercoles": "Miércoles", "jueves": "Jueves", "viernes": "Viernes",
        "sábado": "Sábado", "sabado": "Sábado", "domingo": "Domingo",
    }
    lowered = clean_text(text).lower()
    return unique_join(label for word, label in canonical.items() if re.search(rf"\b{word}\b", lowered))


def extract_locations(text: str, terms: list[str]) -> list[str]:
    haystack = slugify(f"{text} {' '.join(terms)}")
    aliases = {"Lambayaque": "Lambayeque", "Uyacali": "Ucayali", "Ancasch": "Áncash", "Ancash": "Áncash"}
    return unique(
        aliases.get(location, location)
        for location in KNOWN_LOCATIONS
        if slugify(location) in haystack
    )


def parse_promotion(
    source_html: str,
    url: str,
    *,
    post_id: Any = "",
    title: str = "",
    image: str = "",
    terms: list[str] | None = None,
    technical_source: str,
) -> dict[str, Any]:
    soup = BeautifulSoup(source_html, "html.parser")
    vigencia_heading = next(
        (h for h in soup.find_all(re.compile(r"^h[1-6]$")) if slugify(h.get_text(" ", strip=True)) == "vigencia"),
        None,
    )
    merchant = nearest_heading_before(vigencia_heading, "h3")
    benefit = nearest_heading_before(vigencia_heading, "h2")
    if not merchant:
        merchant = html_text(title)
    if not benefit:
        candidates = [clean_text(h.get_text(" ", strip=True)) for h in soup.find_all(("h1", "h2"))]
        benefit = next((x for x in candidates if x and slugify(x) != slugify(merchant)), "")

    vigencia = section_from_heading(soup, "Vigencia")
    conditions = section_from_heading(soup, "Condiciones")
    restrictions = section_from_heading(soup, "Restricciones")
    full_text = clean_text(soup.get_text("\n", strip=True))

    description = ""
    if vigencia_heading:
        benefit_heading = vigencia_heading.find_previous("h2")
        if benefit_heading:
            candidates: list[str] = []
            for node in benefit_heading.next_elements:
                if node is vigencia_heading:
                    break
                if isinstance(node, Tag) and node.name == "p":
                    value = clean_text(node.get_text(" ", strip=True))
                    if value and len(value) <= 800 and not re.match(
                        r"^(vence en|compartir promoción)", value, re.I
                    ):
                        candidates.append(value)
            description = next((x for x in unique(candidates) if x != benefit), "")

    combined = clean_text("\n".join((merchant, benefit, description, vigencia, conditions, restrictions)))
    dates = visible_dates(f"{vigencia}\n{conditions}")
    start = min(dates) if len(dates) >= 2 else None
    end = max(dates) if dates else None
    promo_price, regular_price = money_values(combined)
    term_values = terms or []
    locations = extract_locations(combined, term_values)
    links = unique(urljoin(BASE_URL, a.get("href")) for a in soup.select("a[href]"))
    if not image:
        image_tag = soup.select_one("article img[src], main img[src], meta[property='og:image']")
        if image_tag:
            image = clean_text(image_tag.get("src") or image_tag.get("content"))
    stable_id = clean_text(post_id) or hashlib.sha1(url.encode("utf-8")).hexdigest()[:16]

    category_terms = [
        item for item in term_values
        if slugify(item) not in {slugify(x) for x in locations}
    ]
    return {
        "Banco": "Tarjeta Cencosud",
        "Tipo": unique_join(category_terms) or "Promociones",
        "Nombre": merchant,
        "Detalle de la promoción": description,
        "Restricciones": restrictions,
        "Términos y condiciones": conditions,
        "Ubicación": unique_join(locations),
        "Link": url,
        "Categorías": unique_join(category_terms),
        "Establecimiento": merchant,
        "ID promoción": stable_id,
        "Descuento (%)": percentage(combined),
        "Texto del beneficio": benefit,
        "Precio promocional": promo_price,
        "Precio regular": regular_price,
        "Cuotas sin intereses": installments(combined),
        "Fecha de inicio": start.isoformat() if start else "",
        "Fecha de fin": end.isoformat() if end else "",
        "Estado de vigencia": validity(start, end),
        "Días válidos": valid_days(combined),
        "Tarjetas válidas": "Tarjeta de Crédito Cencosud",
        "Regiones / zonas": unique_join(locations),
        "Comercio": merchant,
        "Código de cupón": coupon_code(combined),
        "Enlaces relacionados": unique_join(links),
        "Imagen": urljoin(BASE_URL, image),
        "Fuente técnica": technical_source,
        "Estado de extracción": "Detalle completo" if conditions or vigencia else "Detalle parcial",
        "Fecha de extracción (UTC)": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
    }


def row_from_post(post: dict[str, Any]) -> dict[str, Any]:
    content = post.get("content", {}).get("rendered", "")
    title = post.get("title", {}).get("rendered", "")
    return parse_promotion(
        content,
        clean_text(post.get("link")),
        post_id=post.get("id"),
        title=title,
        image=featured_image(post),
        terms=embedded_terms(post),
        technical_source="WordPress REST API",
    )


def fetch_html_rows(
    session: requests.Session,
    urls: list[str],
    workers: int,
) -> list[dict[str, Any]]:
    def download(url: str) -> dict[str, Any]:
        response = checked_get(session, url)
        return parse_promotion(response.text, url, technical_source="WordPress HTML")

    rows: list[dict[str, Any]] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(download, url): url for url in urls}
        for count, future in enumerate(as_completed(futures), 1):
            try:
                rows.append(future.result())
            except Exception as exc:
                url = futures[future]
                rows.append({
                    **{column: "" for column in COLUMNS},
                    "Banco": "Tarjeta Cencosud",
                    "Link": url,
                    "Fuente técnica": "WordPress HTML",
                    "Estado de extracción": f"Error: {type(exc).__name__}: {exc}",
                    "Fecha de extracción (UTC)": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
                })
            if count % 25 == 0 or count == len(futures):
                print(f"  Fichas HTML: {count}/{len(futures)}", flush=True)
    return rows


def deduplicate(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    for row in rows:
        key = clean_text(row.get("Link")) or clean_text(row.get("ID promoción"))
        if key and key not in output:
            output[key] = row
    return list(output.values())


def save_excel(rows: list[dict[str, Any]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    frame = pd.DataFrame(rows, columns=COLUMNS)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Promociones")
        sheet = writer.book["Promociones"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 34
        fill = PatternFill("solid", fgColor="6D2077")
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wide = {
            "Nombre": 30, "Detalle de la promoción": 48, "Restricciones": 55,
            "Términos y condiciones": 75, "Ubicación": 28, "Link": 48,
            "Texto del beneficio": 38, "Enlaces relacionados": 52, "Imagen": 48,
            "Estado de extracción": 38,
        }
        for index, name in enumerate(COLUMNS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = wide.get(name, 19)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Promociones Cencosud sin Selenium")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Archivo Excel de salida")
    parser.add_argument("--fuente", choices=("auto", "api", "html"), default="auto")
    parser.add_argument("--workers", type=int, default=8, help="Descargas HTML simultáneas")
    parser.add_argument("--limit", type=int, default=0, help="Máximo de promociones; 0 = todas")
    parser.add_argument("--solo-vigentes", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = arguments()
    session = make_session()
    rows: list[dict[str, Any]] = []

    if args.fuente in ("auto", "api"):
        print("Buscando WordPress REST API...", flush=True)
        rest_base = find_rest_base(session)
        if rest_base:
            print(f"API encontrada: /wp-json/wp/v2/{rest_base}", flush=True)
            try:
                posts = fetch_rest_posts(session, rest_base, args.limit)
                rows = [row_from_post(post) for post in posts]
                # Si los campos personalizados de WordPress no están incluidos
                # en REST, se completa únicamente esas filas desde su ficha.
                incomplete_urls = [
                    clean_text(row.get("Link"))
                    for row in rows
                    if row.get("Estado de extracción") == "Detalle parcial"
                    and clean_text(row.get("Link"))
                ]
                if incomplete_urls:
                    print(
                        f"Completando {len(incomplete_urls)} fichas no expuestas por REST...",
                        flush=True,
                    )
                    completed = fetch_html_rows(session, incomplete_urls, args.workers)
                    completed_by_url = {clean_text(row.get("Link")): row for row in completed}
                    rows = [completed_by_url.get(clean_text(row.get("Link")), row) for row in rows]
            except Exception as exc:
                if args.fuente == "api":
                    raise
                print(f"REST no disponible ({type(exc).__name__}); usando HTML.", flush=True)
        elif args.fuente == "api":
            raise RuntimeError("El sitio no expone el post type de promociones en WordPress REST.")

    if not rows:
        print("Descubriendo fichas públicas sin navegador...", flush=True)
        urls = discover_from_sitemap(session)
        if not urls:
            urls = discover_from_catalog(session)
        if args.limit:
            urls = urls[: args.limit]
        if not urls:
            raise RuntimeError("No se encontraron URLs de promociones en el sitemap ni en el catálogo.")
        print(f"Promociones descubiertas: {len(urls)}", flush=True)
        rows = fetch_html_rows(session, urls, args.workers)

    rows = deduplicate(rows)
    if args.solo_vigentes:
        rows = [row for row in rows if row.get("Estado de vigencia") == "Vigente"]
    rows.sort(key=lambda row: (clean_text(row.get("Nombre")), clean_text(row.get("Texto del beneficio"))))
    output = Path(args.output).expanduser().resolve()
    save_excel(rows, output)
    print(f"Filas exportadas: {len(rows)}", flush=True)
    print(f"Excel generado: {output}", flush=True)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nEjecución cancelada.")
        raise SystemExit(130)
    except Exception as error:
        print(f"ERROR: {error}")
        raise SystemExit(1)
