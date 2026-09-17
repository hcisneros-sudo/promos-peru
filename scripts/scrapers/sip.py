#!/usr/bin/env python3
"""Descarga el catálogo público de SIP Beneficios y genera un Excel.

No usa Selenium, no inicia sesión y no transmite DNI ni correo electrónico.

Instalación:
    pip install requests openpyxl

Uso:
    python SIP_Promociones_API.py
    python SIP_Promociones_API.py --salida SIP_Promociones.xlsx --workers 10
"""

from __future__ import annotations

import argparse
import html
import json
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


PORTAL_URL = "https://beneficios.sip.pe/"
DEFAULT_API_BASE = "https://apis-saas-prd.tarjetaoh.com.pe"
IMAGE_BASE = "https://bucket.sip.pe"
TIMEOUT = 45


def build_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=5,
        connect=5,
        read=5,
        backoff_factor=1,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET", "POST"}),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_connections=20, pool_maxsize=20)
    session.mount("https://", adapter)
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/152.0.0.0 Safari/537.36"
            ),
            "Origin": PORTAL_URL.rstrip("/"),
            "Referer": PORTAL_URL,
            "Accept": "application/json, text/plain, */*",
        }
    )
    return session


def first_match(pattern: str, text: str, default: str = "") -> str:
    found = re.search(pattern, text)
    return found.group(1) if found else default


def discover_public_config() -> dict[str, str]:
    """Lee la configuración pública que el propio portal entrega al navegador."""
    session = build_session()
    response = session.get(PORTAL_URL, timeout=TIMEOUT)
    response.raise_for_status()

    assets = re.findall(
        r'(?:src|href)=["\']([^"\']+\.js(?:\?[^"\']*)?)["\']', response.text
    )
    assets = [urljoin(PORTAL_URL, asset) for asset in assets]
    assets = [
        asset
        for asset in assets
        if urlparse(asset).hostname == urlparse(PORTAL_URL).hostname
    ]
    assets = list(dict.fromkeys(assets))
    if not assets:
        raise RuntimeError("No se encontraron los archivos JavaScript del portal SIP.")

    combined = ""
    for index, asset in enumerate(assets, start=1):
        try:
            js = session.get(asset, timeout=TIMEOUT)
            js.raise_for_status()
        except requests.RequestException:
            continue
        if "client_id_public" in js.text or "base_url" in js.text:
            combined += "\n" + js.text
        if "client_id_public" in combined and "client_secret_public" in combined:
            break

    client_id = first_match(r'client_id_public\s*:\s*["\']([^"\']+)', combined)
    client_secret = first_match(
        r'client_secret_public\s*:\s*["\']([^"\']+)', combined
    )
    api_base = first_match(
        r'(?<!wp_)\bbase_url\s*:\s*["\'](https://[^"\']+)',
        combined,
        DEFAULT_API_BASE,
    )
    version = first_match(r'version\s*:\s*["\']([0-9.]+)', combined, "2.0.1")

    if not client_id or not client_secret:
        raise RuntimeError(
            "SIP cambió la ubicación de su configuración pública. "
            "Debe actualizarse el detector de client_id_public/client_secret_public."
        )

    return {
        "client_id": client_id,
        "client_secret": client_secret,
        "api_base": api_base.rstrip("/"),
        "version": version,
    }


class SipApi:
    def __init__(self, config: dict[str, str]) -> None:
        self.client_id = config["client_id"]
        self.client_secret = config["client_secret"]
        self.api_base = config["api_base"]
        self.version = config["version"]
        self._token = ""
        self._expires_at = 0.0
        self._lock = threading.Lock()
        self._local = threading.local()

    def _session(self) -> requests.Session:
        if not hasattr(self._local, "session"):
            self._local.session = build_session()
        return self._local.session

    def _technical_headers(self) -> dict[str, str]:
        return {
            "codigo-canal": "7",
            "x-application": "WEB",
            "x-plataform": "WEB",
            "x-version": self.version,
            "x-device-id": "Chrome",
            "x-device-ip": "gcp",
            "x-model": "WEB",
        }

    def _authenticate_locked(self) -> None:
        headers = self._technical_headers()
        headers["x-skip-interceptor"] = "true"
        response = self._session().post(
            f"{self.api_base}/public/auth/v1/accesstoken",
            data={
                "client_id": self.client_id,
                "client_secret": self.client_secret,
                "grant_type": "client_credentials",
            },
            headers=headers,
            timeout=TIMEOUT,
        )
        response.raise_for_status()
        payload = response.json()
        self._token = payload["access_token"]
        lifetime = int(payload.get("expires_in", 299))
        self._expires_at = time.monotonic() + max(30, lifetime - 30)

    def ensure_token(self, force: bool = False) -> None:
        if not force and self._token and time.monotonic() < self._expires_at:
            return
        with self._lock:
            if force or not self._token or time.monotonic() >= self._expires_at:
                self._authenticate_locked()

    def get(self, path: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
        self.ensure_token()
        for attempt in range(2):
            headers = self._technical_headers()
            headers.update(
                {
                    "Authorization": f"Bearer {self._token}",
                    "x-api-key": self.client_id,
                }
            )
            response = self._session().get(
                f"{self.api_base}{path}",
                params=params,
                headers=headers,
                timeout=TIMEOUT,
            )
            if response.status_code == 401 and attempt == 0:
                self.ensure_token(force=True)
                continue
            response.raise_for_status()
            return response.json()
        raise RuntimeError(f"No se pudo consultar {path}")


def unwrap_data(payload: dict[str, Any]) -> Any:
    data = payload.get("data", payload)
    if isinstance(data, dict) and "data" in data:
        return data["data"]
    return data


def fetch_all_promotions(api: SipApi, page_size: int) -> list[dict[str, Any]]:
    items: dict[int, dict[str, Any]] = {}
    page = 0
    total_pages = 1
    while page < total_pages:
        payload = api.get(
            "/ohbpromocion/v1/promociones",
            {
                "page": page,
                "size": page_size,
                "esActivo": "true",
                "esCaducado": "false",
            },
        )
        data = payload.get("data") or {}
        promotions = data.get("promociones") or []
        total_pages = int(data.get("nroPages") or 1)
        total_items = int(data.get("items") or len(promotions))
        for promotion in promotions:
            if promotion.get("id") is not None:
                items[int(promotion["id"])] = promotion
        print(
            f"  Catálogo: página {page + 1}/{total_pages} "
            f"({len(items)}/{total_items})"
        )
        page += 1
    return list(items.values())


def fetch_category_map(api: SipApi, page_size: int) -> dict[int, list[str]]:
    payload = api.get("/ohbcategoria/v1/categorias")
    categories = unwrap_data(payload) or []
    mapping: dict[int, set[str]] = {}

    for position, category in enumerate(categories, start=1):
        slug = category.get("slug")
        name = category.get("nombre") or slug
        if not slug:
            continue
        page = 0
        total_pages = 1
        while page < total_pages:
            result = api.get(
                f"/ohbpromocion/v1/promociones/categoria/{slug}",
                {"page": page, "size": page_size, "esActivo": "true"},
            )
            data = result.get("data") or {}
            total_pages = int(data.get("nroPages") or 1)
            for promotion in data.get("promociones") or []:
                promo_id = promotion.get("id")
                if promo_id is not None:
                    mapping.setdefault(int(promo_id), set()).add(str(name))
            page += 1
        print(f"  Categorías: {position}/{len(categories)} - {name}")

    return {key: sorted(value) for key, value in mapping.items()}


def fetch_detail(api: SipApi, promotion: dict[str, Any]) -> dict[str, Any]:
    slug = promotion.get("slug")
    promo_id = promotion.get("id")
    identifier = slug or str(promo_id)
    try:
        payload = api.get(f"/ohbpromocion/v1/promociones/slug/{identifier}")
    except requests.HTTPError as exc:
        if exc.response is None or exc.response.status_code not in (400, 404):
            raise
        payload = api.get(f"/ohbpromocion/v1/promociones/{promo_id}")
    detail = unwrap_data(payload)
    return detail if isinstance(detail, dict) else {}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    text = html.unescape(str(value))
    text = text.replace("\\n", "\n").replace("\r", "\n")
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\[([^]]+)]\([^)]+\)", r"\1", text)
    text = re.sub(r"[*_#`]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def image_url(folder: str, filename: Any) -> str:
    if not filename:
        return ""
    value = str(filename)
    return value if value.startswith("http") else f"{IMAGE_BASE}/{folder}/{value}"


def build_row(
    promotion: dict[str, Any],
    detail: dict[str, Any],
    category_map: dict[int, list[str]],
    extracted_at: str,
) -> dict[str, Any]:
    promo_id = int(detail.get("id") or promotion.get("id"))
    slug = detail.get("slug") or promotion.get("slug") or str(promo_id)
    merchant = promotion.get("comercioNombre") or ""
    return {
        "Banco / marca": "SIP",
        "ID": promo_id,
        "Categorías": ", ".join(category_map.get(promo_id, [])),
        "Comercio": clean_text(merchant),
        "Promoción": clean_text(detail.get("nombre") or promotion.get("sumilla")),
        "Beneficio": clean_text(detail.get("info") or promotion.get("info")),
        "Resumen": clean_text(detail.get("sumilla") or promotion.get("sumilla")),
        "Descripción": clean_text(detail.get("descripcion")),
        "Características": clean_text(detail.get("caracteristicas")),
        "Tarjeta / medio de pago": clean_text(
            detail.get("tarjeta") or promotion.get("tarjeta")
        ),
        "Fecha de inicio": clean_text(
            detail.get("fechaInicio") or promotion.get("fechaInicio")
        ),
        "Fecha de término": clean_text(
            detail.get("fechaTermino") or promotion.get("fechaTermino")
        ),
        "Disponibilidad": clean_text(detail.get("disponibilidad")),
        "Términos y condiciones": clean_text(detail.get("terminosCondiciones")),
        "Estado": clean_text(detail.get("estado") or "ACTIVO"),
        "Texto del enlace": clean_text(detail.get("textoEnlace")),
        "Enlace externo": clean_text(detail.get("urlEnlace")),
        "URL de la promoción": urljoin(PORTAL_URL, f"promociones/{slug}"),
        "Logo del comercio": image_url("Comercio/icon", promotion.get("comercioLogo")),
        "Imagen web": image_url("promociones/web", detail.get("imagenWeb")),
        "Imagen móvil": image_url("promociones/mobile", detail.get("imagenMobile")),
        "Miniatura": image_url(
            "promociones/thumbnail",
            detail.get("thumbnail") or promotion.get("thumbnail"),
        ),
        "Fecha de extracción": extracted_at,
    }


def autosize_sheet(sheet) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            length = min(len(str(cell.value or "")), 60)
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = max(12, width + 2)


def save_excel(rows: list[dict[str, Any]], output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Promociones"

    headers = list(rows[0].keys()) if rows else ["Resultado"]
    sheet.append(headers)
    if not rows:
        sheet.append(["No se encontraron promociones activas."])
    else:
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="00B1FF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize_sheet(sheet)

    summary = workbook.create_sheet("Resumen")
    summary.append(["Indicador", "Valor"])
    summary.append(["Fuente", PORTAL_URL])
    summary.append(["Promociones descargadas", len(rows)])
    summary.append(["Fecha de extracción", datetime.now().strftime("%d/%m/%Y %H:%M:%S")])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
    autosize_sheet(summary)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Descarga las promociones públicas de SIP Beneficios a Excel."
    )
    parser.add_argument("--salida", default="SIP_Promociones.xlsx")
    parser.add_argument("--workers", type=int, default=10)
    parser.add_argument("--tamano-pagina", type=int, default=100)
    parser.add_argument(
        "--limite",
        type=int,
        default=0,
        help="Limita las promociones procesadas; 0 descarga todas.",
    )
    parser.add_argument(
        "--sin-detalle",
        action="store_true",
        help="Omite las consultas de detalle para una ejecución más rápida.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.workers < 1 or args.workers > 20:
        raise SystemExit("--workers debe estar entre 1 y 20.")

    print("Detectando la configuración pública de SIP Beneficios...")
    config = discover_public_config()
    api = SipApi(config)
    api.ensure_token()

    print("Descargando catálogo completo mediante API...")
    promotions = fetch_all_promotions(api, max(10, args.tamano_pagina))
    if args.limite > 0:
        promotions = promotions[: args.limite]

    print("Relacionando promociones con categorías...")
    category_map = fetch_category_map(api, max(20, args.tamano_pagina))

    details: dict[int, dict[str, Any]] = {}
    if not args.sin_detalle and promotions:
        print(f"Descargando {len(promotions)} detalles con {args.workers} workers...")
        with ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = {
                executor.submit(fetch_detail, api, promotion): promotion
                for promotion in promotions
            }
            for position, future in enumerate(as_completed(futures), start=1):
                promotion = futures[future]
                promo_id = int(promotion["id"])
                try:
                    details[promo_id] = future.result()
                except Exception as exc:  # conserva el catálogo aunque falle un detalle
                    print(f"  Aviso: detalle {promo_id} no disponible: {exc}")
                    details[promo_id] = {}
                if position % 20 == 0 or position == len(promotions):
                    print(f"  Detalles: {position}/{len(promotions)}")

    extracted_at = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    rows = [
        build_row(
            promotion,
            details.get(int(promotion["id"]), {}),
            category_map,
            extracted_at,
        )
        for promotion in promotions
    ]
    rows.sort(key=lambda row: (row["Comercio"].lower(), row["Promoción"].lower()))

    output = Path(args.salida).expanduser().resolve()
    save_excel(rows, output)
    print(f"Listo: {len(rows)} promociones guardadas en:")
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
