"""Scraper de promociones de Banco Falabella Peru.

La web usa Next.js y entrega los datos de promociones dentro del HTML como
React Flight data. Este programa lee esos datos estructurados sin Selenium,
visita los detalles en paralelo y genera un Excel.

Uso:
    python Banco_Falabella_Promociones.py
    python Banco_Falabella_Promociones.py --limit 10
    python Banco_Falabella_Promociones.py --no-details

Dependencias:
    pip install pandas openpyxl
"""

from __future__ import annotations

import argparse
import html as html_lib
import json
import random
import re
import socket
import ssl
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_URL = "https://www.bancofalabella.pe"
LIST_URL = f"{BASE_URL}/promociones/todos"
DEFAULT_OUTPUT = "promociones_banco_falabella.xlsx"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/152.0.0.0 Safari/537.36"
)

# Las primeras columnas conservan la logica del ejemplo de Banco Ripley.
COLUMNS = [
    "Banco",
    "Tipo",
    "Nombre",
    "Detalle de la promoción",
    "Restricciones",
    "Términos y condiciones",
    "Ubicación",
    "Link",
    "Subcategorías",
    "Descuento (%)",
    "Texto del beneficio",
    "Fecha de inicio",
    "Fecha de fin",
    "Días válidos",
    "Tarjetas válidas",
    "Regiones",
    "Ciudades",
    "Modalidad",
    "Comercio",
    "Cómo acceder / detalle",
    "Código de cupón",
    "Enlace informativo",
    "Imagen",
    "Logo",
    "Es nueva",
    "Destacada",
    "Estado de extracción",
    "Fecha de extracción (UTC)",
]


def clean_text(value: Any) -> str:
    """Limpia espacios y entidades HTML sin destruir saltos de linea."""
    if value is None:
        return ""
    text = html_lib.unescape(str(value)).replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def join_values(value: Any, separator: str = " | ") -> str:
    if value in (None, ""):
        return ""
    if not isinstance(value, (list, tuple, set)):
        return clean_text(value)
    values = []
    for item in value:
        text = clean_text(item)
        if text and text not in values:
            values.append(text)
    return separator.join(values)


def absolute_url(value: Any) -> str:
    value = clean_text(value)
    if not value:
        return ""
    if value.startswith("//"):
        return "https:" + value
    return urljoin(BASE_URL, value)


def fetch_html(url: str, attempts: int = 4, timeout: int = 60) -> str:
    """Descarga una pagina con reintentos y espera exponencial."""
    headers = {
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "es-PE,es;q=0.9,en;q=0.7",
        "Cache-Control": "no-cache",
    }
    context = ssl.create_default_context()
    last_error: Exception | None = None

    for attempt in range(1, attempts + 1):
        try:
            request = Request(url, headers=headers)
            with urlopen(request, timeout=timeout, context=context) as response:
                raw = response.read()
                charset = response.headers.get_content_charset() or "utf-8"
                return raw.decode(charset, errors="replace")
        except (HTTPError, URLError, TimeoutError, socket.timeout) as exc:
            last_error = exc
            if attempt == attempts:
                break
            time.sleep((1.4 ** attempt) + random.uniform(0.1, 0.6))

    raise RuntimeError(f"No se pudo descargar {url}: {last_error}")


def extract_flight_text(page_html: str) -> str:
    """Decodifica los strings enviados mediante self.__next_f.push(...)."""
    parts: list[str] = []
    pattern = re.compile(r"self\.__next_f\.push\((\[.*?\])\)</script>", re.S)
    for match in pattern.finditer(page_html):
        try:
            payload = json.loads(match.group(1))
        except json.JSONDecodeError:
            continue
        if len(payload) > 1 and isinstance(payload[1], str):
            parts.append(payload[1])
    if not parts:
        raise ValueError("La pagina no contiene datos Next.js reconocibles")
    return "".join(parts)


def json_values_after_key(flight_text: str, key: str) -> list[Any]:
    """Extrae todos los valores JSON validos que aparecen despues de una clave."""
    decoder = json.JSONDecoder()
    marker = json.dumps(key, ensure_ascii=False) + ":"
    values: list[Any] = []
    start = 0
    while True:
        index = flight_text.find(marker, start)
        if index < 0:
            break
        value_start = index + len(marker)
        try:
            value, consumed = decoder.raw_decode(flight_text[value_start:])
            values.append(value)
            start = value_start + consumed
        except json.JSONDecodeError:
            start = value_start
    return values


def extract_text_records(flight_text: str) -> dict[str, str]:
    """Extrae registros de texto React Flight del tipo 80:Tcbc,<texto>."""
    data = flight_text.encode("utf-8")
    records: dict[str, str] = {}
    pattern = re.compile(rb"(?:^|\n)([0-9a-z]+):T([0-9a-f]+),")
    for match in pattern.finditer(data):
        record_id = match.group(1).decode("ascii")
        byte_length = int(match.group(2), 16)
        begin = match.end()
        end = begin + byte_length
        if end <= len(data):
            records[record_id] = data[begin:end].decode("utf-8", errors="replace")
    return records


def resolve_text_reference(value: Any, records: dict[str, str]) -> Any:
    if isinstance(value, str) and re.fullmatch(r"\$[0-9a-z]+", value):
        return records.get(value[1:], value)
    return value


def rich_text_to_plain(value: Any) -> str:
    """Convierte un documento Rich Text de Contentful en texto legible."""
    if not isinstance(value, dict):
        return clean_text(value)

    blocks: list[str] = []

    def walk(node: Any, depth: int = 0) -> str:
        if isinstance(node, list):
            return "".join(walk(item, depth) for item in node)
        if not isinstance(node, dict):
            return ""
        node_type = node.get("nodeType", "")
        if node_type == "text":
            return str(node.get("value", ""))
        content = walk(node.get("content", []), depth + 1)
        content = re.sub(r"[ \t]+", " ", content).strip()
        if node_type == "list-item" and content:
            blocks.append("• " + content)
            return ""
        if node_type in {"paragraph", "heading-1", "heading-2", "heading-3"} and content:
            blocks.append(content)
            return ""
        return content

    walk(value)
    return clean_text("\n".join(blocks))


def locations_to_text(locations: Any) -> str:
    if not locations:
        return ""
    if not isinstance(locations, list):
        locations = [locations]
    output: list[str] = []
    preferred = (
        "name", "title", "commerceName", "address", "direction", "direccion",
        "district", "city", "region", "schedule", "phone",
    )
    for location in locations:
        if isinstance(location, dict):
            parts = []
            for key in preferred:
                if location.get(key) not in (None, "", []):
                    text = join_values(location[key], ", ")
                    if text and text not in parts:
                        parts.append(text)
            if not parts:
                parts = [clean_text(v) for v in location.values() if isinstance(v, (str, int, float))]
            line = " — ".join(filter(None, parts))
        else:
            line = clean_text(location)
        if line and line not in output:
            output.append(line)
    return " | ".join(output)


def extract_restrictions(legal_text: str) -> str:
    """Crea un resumen literal de clausulas restrictivas; el texto legal queda completo aparte."""
    if not legal_text:
        return ""
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÑ0-9])", legal_text)
    indicators = (
        "no válido", "no valida", "no aplica", "no acumulable", "excepto",
        "máximo", "maximo", "mínimo", "minimo", "stock", "sujeto a",
        "exclusivo", "solo ", "sólo ", "restricción", "recargo",
    )
    selected = [s.strip() for s in sentences if any(x in s.lower() for x in indicators)]
    return clean_text(" ".join(selected))


def iso_to_date(value: Any) -> str:
    value = clean_text(value)
    if not value:
        return ""
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return value


def discount_percent(card: dict[str, Any], detail: dict[str, Any] | None = None) -> Any:
    """Obtiene un porcentaje solo cuando la tarjeta realmente muestra el signo %."""
    detail = detail or {}
    text = " ".join(
        clean_text(source.get(key))
        for source in (detail, card)
        for key in ("topDiscountText", "centerDiscountText", "bottomDiscountText")
        if clean_text(source.get(key))
    )
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*%", text)
    if not match:
        return ""
    number = match.group(1).replace(",", ".")
    return float(number) if "." in number else int(number)


def extract_list_items(page_html: str) -> list[dict[str, Any]]:
    flight = extract_flight_text(page_html)
    candidates = [x for x in json_values_after_key(flight, "benefitCardsData") if isinstance(x, list)]
    if not candidates:
        raise ValueError("No se encontro benefitCardsData en la pagina de promociones")
    # Puede haber copias parciales o duplicadas. La lista mayor es la completa.
    items = max(candidates, key=len)
    unique: dict[str, dict[str, Any]] = {}
    for item in items:
        if not isinstance(item, dict):
            continue
        card = item.get("benefitCard") or {}
        link = absolute_url(card.get("linkUrl"))
        if "/promociones/detalle/" in link:
            unique.setdefault(link, item)
    return list(unique.values())


def extract_detail(page_html: str, expected_slug: str) -> dict[str, Any]:
    flight = extract_flight_text(page_html)
    records = extract_text_records(flight)
    candidates = [x for x in json_values_after_key(flight, "benefitData") if isinstance(x, dict)]
    if not candidates:
        raise ValueError("No se encontro benefitData")
    detail = next(
        (x for x in candidates if x.get("permalink") == expected_slug),
        candidates[0],
    )
    detail["legalText"] = clean_text(resolve_text_reference(detail.get("legalText"), records))
    return detail


def fallback_row(item: dict[str, Any], error: str = "") -> dict[str, Any]:
    card = item.get("benefitCard") or {}
    link = absolute_url(card.get("linkUrl"))
    discount_text = " ".join(
        clean_text(card.get(k))
        for k in ("topDiscountText", "centerDiscountText", "bottomDiscountText")
        if clean_text(card.get(k))
    )
    return {
        "Banco": "Banco Falabella",
        "Tipo": "",
        "Nombre": clean_text(card.get("title") or item.get("benefitTitle")),
        "Detalle de la promoción": clean_text(card.get("description")),
        "Restricciones": "",
        "Términos y condiciones": "",
        "Ubicación": join_values(item.get("cities") if isinstance(item.get("cities"), list) else item.get("region")),
        "Link": link,
        "Subcategorías": "",
        "Descuento (%)": discount_percent(card),
        "Texto del beneficio": discount_text,
        "Fecha de inicio": iso_to_date(card.get("initDate")),
        "Fecha de fin": iso_to_date(card.get("endDate") or item.get("limitDate")),
        "Días válidos": join_values(card.get("discountDays")),
        "Tarjetas válidas": join_values(item.get("creditCards")),
        "Regiones": join_values(item.get("region")),
        "Ciudades": join_values(item.get("cities") if isinstance(item.get("cities"), list) else []),
        "Modalidad": "",
        "Comercio": clean_text(item.get("benefitTitle")),
        "Cómo acceder / detalle": "",
        "Código de cupón": "",
        "Enlace informativo": "",
        "Imagen": absolute_url(card.get("imageCard")),
        "Logo": absolute_url(card.get("logoCard")),
        "Es nueva": "Sí" if card.get("isNew") else "No",
        "Destacada": "Sí" if item.get("highlighted") else "No",
        "Estado de extracción": "OK (listado)" if not error else f"Detalle no disponible: {error}",
        "Fecha de extracción (UTC)": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
    }


def build_row(item: dict[str, Any], detail: dict[str, Any]) -> dict[str, Any]:
    row = fallback_row(item)
    card = item.get("benefitCard") or {}
    categories = detail.get("relatedCategory") or []
    legal = clean_text(detail.get("legalText"))
    locations = locations_to_text(detail.get("locations"))
    regions = detail.get("region") or item.get("region") or []
    row.update({
        "Tipo": clean_text(categories[0]) if categories else "",
        "Nombre": clean_text(detail.get("cardTitle") or detail.get("benefitTitle") or card.get("title")),
        "Detalle de la promoción": clean_text(detail.get("cardDescription") or card.get("description")),
        "Restricciones": extract_restrictions(legal),
        "Términos y condiciones": legal,
        "Ubicación": locations or join_values(regions),
        "Subcategorías": join_values(categories),
        "Descuento (%)": discount_percent(card, detail),
        "Fecha de inicio": iso_to_date(detail.get("initDate") or card.get("initDate")),
        "Fecha de fin": iso_to_date(detail.get("endDate") or card.get("endDate")),
        "Días válidos": join_values(detail.get("discountDays") or card.get("discountDays")),
        "Tarjetas válidas": join_values(detail.get("creditCards") or item.get("creditCards")),
        "Regiones": join_values(regions),
        "Ciudades": join_values(item.get("cities") if isinstance(item.get("cities"), list) else []),
        "Modalidad": join_values(detail.get("benefitsMode")),
        "Comercio": clean_text(detail.get("commerceName") or item.get("benefitTitle")),
        "Cómo acceder / detalle": rich_text_to_plain(detail.get("detailBanner1")),
        "Código de cupón": clean_text(detail.get("couponCode")),
        "Enlace informativo": absolute_url(detail.get("commerceInfoUrl") or detail.get("urlCta")),
        "Imagen": absolute_url(detail.get("cardImage") or card.get("imageCard")),
        "Logo": absolute_url(detail.get("cardLogo") or card.get("logoCard")),
        "Es nueva": "Sí" if detail.get("isNewCard", card.get("isNew")) else "No",
        "Destacada": "Sí" if detail.get("highlighted", item.get("highlighted")) else "No",
        "Estado de extracción": "OK",
    })
    return row


def scrape_one_detail(item: dict[str, Any]) -> dict[str, Any]:
    card = item.get("benefitCard") or {}
    link = absolute_url(card.get("linkUrl"))
    slug = urlparse(link).path.rstrip("/").split("/")[-1]
    try:
        detail = extract_detail(fetch_html(link), slug)
        return build_row(item, detail)
    except Exception as exc:  # La fila del listado nunca se pierde.
        return fallback_row(item, f"{type(exc).__name__}: {exc}")


def format_excel(path: Path, row_count: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Promociones"]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    header_fill = PatternFill("solid", fgColor="4F7F36")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "Banco": 20, "Tipo": 20, "Nombre": 38, "Detalle de la promoción": 48,
        "Restricciones": 55, "Términos y condiciones": 75, "Ubicación": 45,
        "Link": 48, "Subcategorías": 30, "Descuento (%)": 15,
        "Texto del beneficio": 25, "Fecha de inicio": 15, "Fecha de fin": 15,
        "Días válidos": 35, "Tarjetas válidas": 45, "Regiones": 45,
        "Ciudades": 35, "Modalidad": 20, "Comercio": 35,
        "Cómo acceder / detalle": 60, "Código de cupón": 18,
        "Enlace informativo": 45, "Imagen": 45, "Logo": 45,
        "Es nueva": 12, "Destacada": 12, "Estado de extracción": 34,
        "Fecha de extracción (UTC)": 23,
    }
    for index, name in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 20)

    for row in sheet.iter_rows(min_row=2, max_row=row_count + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.sheet_view.showGridLines = False
    workbook.save(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extrae promociones de Banco Falabella a Excel")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Ruta del Excel de salida")
    parser.add_argument("--workers", type=int, default=6, help="Descargas simultaneas (1-12)")
    parser.add_argument("--limit", type=int, default=0, help="Limite de promociones; 0 = todas")
    parser.add_argument("--no-details", action="store_true", help="No visitar fichas; exportacion rapida")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = Path(args.output).expanduser().resolve()
    workers = max(1, min(args.workers, 12))
    started = time.time()

    print(f"Descargando listado: {LIST_URL}")
    items = extract_list_items(fetch_html(LIST_URL))
    if args.limit > 0:
        items = items[: args.limit]
    print(f"Promociones unicas encontradas: {len(items)}")

    if args.no_details:
        rows = [fallback_row(item) for item in items]
    else:
        rows = []
        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(scrape_one_detail, item): item for item in items}
            for number, future in enumerate(as_completed(futures), 1):
                row = future.result()
                rows.append(row)
                print(f"[{number}/{len(items)}] {row['Estado de extracción']}: {row['Nombre']}")

    rows.sort(key=lambda x: (x.get("Tipo", ""), x.get("Nombre", "")))
    frame = pd.DataFrame(rows, columns=COLUMNS)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Promociones", index=False)
    format_excel(output, len(frame))

    errors = sum(not str(row["Estado de extracción"]).startswith("OK") for row in rows)
    print(f"\nExcel generado: {output}")
    print(f"Filas: {len(frame)} | Detalles con error: {errors}")
    print(f"Tiempo total: {time.time() - started:.1f} segundos")


if __name__ == "__main__":
    main()
