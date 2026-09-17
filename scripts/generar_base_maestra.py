from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


MASTER_COLUMNS = [
    "id_registro", "id_promocion", "id_local", "banco", "comercio",
    "nombre_promocion", "descripcion", "categoria", "subcategoria",
    "tipo_beneficio", "porcentaje_descuento", "beneficio_texto", "medio_pago",
    "modalidad", "alcance", "fecha_inicio", "fecha_fin", "dias_validos",
    "estado_vigencia", "restricciones", "terminos_condiciones", "nombre_local", "direccion_original", "direccion_limpia",
    "direccion_geocodificada", "distrito", "provincia", "departamento",
    "latitud", "longitud", "estado_mapa", "precision_ubicacion",
    "confianza_ubicacion", "puntaje_geocodificacion", "fuente_coordenada",
    "correccion_automatica", "regla_correccion", "tipo_revision",
    "prioridad_revision", "accion_sugerida", "resultado_api", "puntaje_revision",
    "url_promocion", "url_imagen", "archivo_origen", "fecha_extraccion",
]

PROMOTION_COLUMNS = [
    "id_promocion", "banco", "comercio", "nombre_promocion", "descripcion",
    "categoria", "subcategoria", "confianza_categoria", "tipo_beneficio",
    "porcentaje_descuento", "beneficio_texto", "medio_pago", "modalidad",
    "alcance", "fecha_inicio", "fecha_fin", "dias_validos", "estado_vigencia",
    "restricciones", "terminos_condiciones", "ubicacion_original",
    "url_promocion", "url_imagen", "archivo_origen", "fecha_extraccion",
]

PENDING_COLUMNS = [
    "id_local", "id_promocion", "banco", "comercio", "nombre_promocion",
    "nombre_local", "direccion_original", "direccion_limpia", "distrito",
    "provincia", "departamento", "estado_mapa", "tipo_revision",
    "prioridad_revision", "accion_sugerida", "resultado_api", "puntaje_revision",
    "archivo_origen",
]

LOCATION_COLUMNS = [
    "id_local", "id_promocion", "nombre_local", "direccion_original",
    "direccion_limpia", "direccion_geocodificada", "distrito", "provincia",
    "departamento", "latitud", "longitud", "estado_mapa", "precision_ubicacion",
    "confianza_ubicacion", "puntaje_geocodificacion", "fuente_coordenada",
    "correccion_automatica", "regla_correccion", "archivo_origen",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=15, bold=True, color="1F1F1F")
BODY_FONT = Font(name="Arial", size=10, color="1F1F1F")
LIGHT_BLUE = PatternFill("solid", fgColor="D9EAF7")
LIGHT_GREEN = PatternFill("solid", fgColor="E2F0D9")
LIGHT_RED = PatternFill("solid", fgColor="FCE4D6")
LIGHT_AMBER = PatternFill("solid", fgColor="FFF2CC")
THIN_GRAY = Side(style="thin", color="D9E1F2")


def read_json(path: Path, required: bool = True):
    if not path.exists():
        if required:
            raise FileNotFoundError(f"No se encontró el archivo requerido: {path}")
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def save_json(path: Path, value, compact: bool = False):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, separators=(",", ":") if compact else None, indent=None if compact else 2),
        encoding="utf-8",
    )


def fold(value: str) -> str:
    return "".join(
        char for char in unicodedata.normalize("NFKD", str(value or "").lower())
        if not unicodedata.combining(char)
    )


def address_key(value: str) -> str:
    text = fold(value)
    for old, new in {
        "prologacion": "prolongacion", "benabides": "benavides",
        "quinonez": "quinones", "lambayaque": "lambayeque",
    }.items():
        text = text.replace(old, new)
    text = re.sub(r"\b(?:avenida|av|jiron|jr|calle|carretera)\b", " ", text)
    text = re.sub(r"\b(?:int(?:erior)?|tda|tienda|local|nivel|piso|oficina|stand|lc)\b.*$", " ", text)
    text = re.sub(r"\b(?:nro|numero|no)\b", " ", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def valid_coordinate(latitude, longitude) -> bool:
    try:
        latitude, longitude = float(latitude), float(longitude)
    except (TypeError, ValueError):
        return False
    return -18.6 <= latitude <= 0.6 and -82.0 <= longitude <= -68.0


def haversine_km(first: tuple[float, float], second: tuple[float, float]) -> float:
    radius = 6371.0
    lat1, lon1, lat2, lon2 = map(math.radians, [first[0], first[1], second[0], second[1]])
    dlat, dlon = lat2 - lat1, lon2 - lon1
    value = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * radius * math.asin(math.sqrt(value))


def inherit_coordinates(locations: list[dict]) -> int:
    groups = defaultdict(list)
    for location in locations:
        if not valid_coordinate(location.get("latitud"), location.get("longitud")):
            continue
        signature = f"{address_key(location.get('direccion_limpia', ''))}|{address_key(location.get('distrito', ''))}"
        if len(signature.replace("|", "")) >= 7:
            groups[signature].append(location)

    reliable = {}
    for signature, rows in groups.items():
        points = [(float(row["latitud"]), float(row["longitud"])) for row in rows]
        center = (
            statistics.median(point[0] for point in points),
            statistics.median(point[1] for point in points),
        )
        if max(haversine_km(center, point) for point in points) <= 0.25:
            addresses = Counter(row.get("direccion_geocodificada", "") for row in rows if row.get("direccion_geocodificada"))
            reliable[signature] = (center, addresses.most_common(1)[0][0] if addresses else "", len(rows))

    inherited = 0
    for location in locations:
        location["correccion_automatica"] = "No"
        location["regla_correccion"] = ""
        if valid_coordinate(location.get("latitud"), location.get("longitud")):
            continue
        signature = f"{address_key(location.get('direccion_limpia', ''))}|{address_key(location.get('distrito', ''))}"
        if signature not in reliable:
            continue
        center, geocoded_address, references = reliable[signature]
        location["latitud"], location["longitud"] = center
        location["direccion_geocodificada"] = geocoded_address
        location["precision_ubicacion"] = "Heredada"
        location["fuente_coordenada"] = "Historial interno / dirección equivalente"
        location["confianza_ubicacion"] = "Media"
        location["correccion_automatica"] = "Sí"
        location["regla_correccion"] = f"Dirección y distrito equivalentes; {references} referencia(s) dentro de 250 m."
        inherited += 1
    return inherited


def review_class(score, reason: str) -> tuple[str, str, str]:
    if "error" in fold(reason):
        return "Error temporal de consulta", "Alta", "Volver a consultar el servicio."
    score = float(score or 0)
    if score >= 0.60:
        return "Coincidencia cercana", "Alta", "Confirmar candidato, distrito y numeración."
    if score >= 0.40:
        return "Dirección parcial", "Media", "Simplificar o completar la dirección y volver a consultar."
    if score >= 0.20:
        return "Coincidencia débil", "Media", "Verificar ciudad, distrito y nombre de la vía."
    return "Sin coincidencia útil", "Baja", "Buscar la dirección oficial del establecimiento."


def build_datasets(input_dir: Path):
    promotion_rows = read_json(input_dir / "promociones.json")
    promotions_by_id = {}
    for promotion in promotion_rows:
        promotions_by_id.setdefault(promotion.get("id_promocion", ""), promotion)
    promotions = list(promotions_by_id.values())
    locations = read_json(input_dir / "locales_geocodificados.json")
    reviews = read_json(input_dir / "revision_geocodificacion.json", required=False)
    quality = read_json(input_dir / "control_calidad.json", required=False)
    duplicates = read_json(input_dir / "duplicados_probables.json", required=False)

    promotion_by_id = {row["id_promocion"]: row for row in promotions}
    review_by_local = defaultdict(list)
    for review in reviews:
        review_by_local[review.get("id_local", "")].append(review)

    inherited = inherit_coordinates(locations)
    master, pending, normalized_locations = [], [], []
    locations_by_promotion = defaultdict(list)
    status_counts = Counter()
    for location in locations:
        locations_by_promotion[location.get("id_promocion", "")].append(location)

    def append_row(promotion: dict, location: dict | None):
        location = location or {}
        review = (review_by_local.get(location.get("id_local", "")) or [{}])[0]
        has_coordinates = valid_coordinate(location.get("latitud"), location.get("longitud"))
        if not location:
            map_status = "SIN LOCAL - ONLINE/NACIONAL" if promotion.get("alcance") in {"Online", "Nacional"} else "SIN LOCAL REGISTRADO"
            review_type = priority = action = ""
        elif location.get("correccion_automatica") == "Sí":
            map_status, review_type, priority, action = "LISTO - HEREDADA", "", "", ""
        elif has_coordinates:
            map_status, review_type, priority, action = "LISTO - COORDENADA", "", "", ""
        elif review:
            map_status = "REVISAR - API"
            review_type, priority, action = review_class(review.get("puntaje"), review.get("motivo", ""))
        else:
            map_status = "REVISAR - ORIGEN"
            address, district = location.get("direccion_limpia", ""), location.get("distrito", "")
            if not re.search(r"\d", address) and not district:
                review_type, priority, action = "Sin dirección verificable", "Alta", "Completar dirección y ciudad."
            elif not re.search(r"\d", address):
                review_type, priority, action = "Solo referencia geográfica", "Media", "Confirmar la dirección exacta."
            else:
                review_type, priority, action = "Dirección excluida por calidad", "Media", "Corregir la estructura y volver a geocodificar."

        status_counts[map_status] += 1
        local_id = location.get("id_local", "")
        record = {
            "id_registro": f"{promotion.get('id_promocion', '')}-{local_id or 'sin-local'}",
            "id_promocion": promotion.get("id_promocion", ""), "id_local": local_id,
            "banco": promotion.get("banco", ""), "comercio": promotion.get("comercio", ""),
            "nombre_promocion": promotion.get("nombre_promocion", ""), "descripcion": promotion.get("descripcion", ""),
            "categoria": promotion.get("categoria", ""), "subcategoria": promotion.get("subcategoria", ""),
            "tipo_beneficio": promotion.get("tipo_beneficio", ""), "porcentaje_descuento": promotion.get("porcentaje_descuento"),
            "beneficio_texto": promotion.get("beneficio_texto", ""), "medio_pago": promotion.get("medio_pago", ""),
            "modalidad": promotion.get("modalidad", ""), "alcance": promotion.get("alcance", ""),
            "fecha_inicio": promotion.get("fecha_inicio", ""), "fecha_fin": promotion.get("fecha_fin", ""),
            "dias_validos": promotion.get("dias_validos", ""), "estado_vigencia": promotion.get("estado_vigencia", ""),
            "restricciones": promotion.get("restricciones", ""), "terminos_condiciones": promotion.get("terminos_condiciones", ""),
            "nombre_local": location.get("nombre_local", ""), "direccion_original": location.get("direccion_original", ""),
            "direccion_limpia": location.get("direccion_limpia", ""), "direccion_geocodificada": location.get("direccion_geocodificada", ""),
            "distrito": location.get("distrito", ""), "provincia": location.get("provincia", ""),
            "departamento": location.get("departamento", ""), "latitud": float(location["latitud"]) if has_coordinates else None,
            "longitud": float(location["longitud"]) if has_coordinates else None, "estado_mapa": map_status,
            "precision_ubicacion": location.get("precision_ubicacion", ""), "confianza_ubicacion": location.get("confianza_ubicacion", ""),
            "puntaje_geocodificacion": location.get("puntaje_geocodificacion"), "fuente_coordenada": location.get("fuente_coordenada", ""),
            "correccion_automatica": location.get("correccion_automatica", "No"), "regla_correccion": location.get("regla_correccion", ""),
            "tipo_revision": review_type, "prioridad_revision": priority, "accion_sugerida": action,
            "resultado_api": review.get("resultado", ""), "puntaje_revision": review.get("puntaje"),
            "url_promocion": promotion.get("url_promocion", ""), "url_imagen": promotion.get("url_imagen", ""),
            "archivo_origen": promotion.get("archivo_origen", location.get("archivo_origen", "")),
            "fecha_extraccion": promotion.get("fecha_extraccion", ""),
        }
        master.append(record)
        if location:
            normalized_locations.append({column: record.get(column) for column in LOCATION_COLUMNS})
        if location and not has_coordinates:
            pending.append({column: record.get(column) for column in PENDING_COLUMNS})

    for promotion in promotions:
        promotion_locations = locations_by_promotion.get(promotion.get("id_promocion", ""), [])
        if promotion_locations:
            for location in promotion_locations:
                append_row(promotion, location)
        else:
            append_row(promotion, None)

    summary = {
        "fecha_generacion": datetime.now().isoformat(timespec="seconds"),
        "total_promociones": len(promotions),
        "filas_promociones_origen": len(promotion_rows),
        "promociones_duplicadas_eliminadas": len(promotion_rows) - len(promotions),
        "total_registros_maestros": len(master),
        "total_locales": len(locations),
        "promociones_sin_local": sum(not locations_by_promotion.get(row.get("id_promocion", "")) for row in promotions),
        "locales_listos_mapa": sum(valid_coordinate(row.get("latitud"), row.get("longitud")) for row in normalized_locations),
        "locales_pendientes": len(pending),
        "coordenadas_heredadas": inherited,
        "casos_revision_api": len(reviews),
        "observaciones_calidad": len(quality),
        "duplicados_probables": len(duplicates),
        "estado_mapa": dict(status_counts),
        "promociones_por_banco": dict(Counter(row.get("banco", "") for row in promotions)),
        "promociones_por_categoria": dict(Counter(row.get("categoria", "") for row in promotions)),
    }
    return promotions, master, normalized_locations, pending, summary


def excel_value(value, column: str):
    if value in ("", None):
        return None
    if column in {"fecha_inicio", "fecha_fin", "fecha_extraccion"}:
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return str(value)
    return value


def add_data_sheet(workbook: Workbook, title: str, rows: list[dict], columns: list[str], table_name: str):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.append(columns)
    for row in rows:
        sheet.append([excel_value(row.get(column), column) for column in columns])
    header = sheet[1]
    for cell in header:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "D2" if len(columns) > 8 else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(columns))}{len(rows)+1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        sheet.add_table(table)

    wide = {"descripcion", "beneficio_texto", "direccion_original", "direccion_limpia", "direccion_geocodificada", "regla_correccion", "accion_sugerida", "resultado_api", "url_promocion", "url_imagen", "terminos_condiciones", "restricciones", "ubicacion_original"}
    medium = {"nombre_promocion", "comercio", "nombre_local", "medio_pago", "archivo_origen"}
    for index, column in enumerate(columns, 1):
        letter = get_column_letter(index)
        width = 42 if column in wide else 28 if column in medium else 18
        sheet.column_dimensions[letter].width = width
        if column in {"latitud", "longitud"}:
            sheet.column_dimensions[letter].width = 14
            for cell in sheet[letter][1:]:
                cell.number_format = "0.000000"
        elif column in {"porcentaje_descuento", "puntaje_geocodificacion", "puntaje_revision"}:
            for cell in sheet[letter][1:]:
                cell.number_format = "0.00"
        elif column in {"fecha_inicio", "fecha_fin", "fecha_extraccion"}:
            sheet.column_dimensions[letter].width = 13
            for cell in sheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    return sheet


def add_summary_sheet(workbook: Workbook, summary: dict):
    sheet = workbook.active
    sheet.title = "Resumen"
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = "Base maestra de promociones bancarias"
    sheet["A2"].font = TITLE_FONT
    sheet["A3"] = f"Generada: {summary['fecha_generacion']}"
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")
    metrics = [
        ("Promociones", summary["total_promociones"]),
        ("Duplicados exactos eliminados", summary["promociones_duplicadas_eliminadas"]),
        ("Registros en base maestra", summary["total_registros_maestros"]),
        ("Locales identificados", summary["total_locales"]),
        ("Puntos listos para mapa", summary["locales_listos_mapa"]),
        ("Locales pendientes", summary["locales_pendientes"]),
        ("Promociones sin local", summary["promociones_sin_local"]),
        ("Coordenadas heredadas", summary["coordenadas_heredadas"]),
    ]
    sheet["A5"], sheet["B5"] = "Indicador", "Cantidad"
    for cell in sheet[5]:
        if cell.column <= 2:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row_index, (label, value) in enumerate(metrics, 6):
        sheet.cell(row_index, 1, label)
        sheet.cell(row_index, 2, value)
        sheet.cell(row_index, 2).number_format = "#,##0"
        if label == "Puntos listos para mapa":
            sheet.cell(row_index, 1).fill = sheet.cell(row_index, 2).fill = LIGHT_GREEN
        elif label == "Locales pendientes":
            sheet.cell(row_index, 1).fill = sheet.cell(row_index, 2).fill = LIGHT_AMBER

    bank_start = 5
    sheet.cell(bank_start, 4, "Banco")
    sheet.cell(bank_start, 5, "Promociones")
    for cell in sheet[bank_start][3:5]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for offset, (bank, count) in enumerate(sorted(summary["promociones_por_banco"].items()), 1):
        sheet.cell(bank_start + offset, 4, bank)
        sheet.cell(bank_start + offset, 5, count)
        sheet.cell(bank_start + offset, 5).number_format = "#,##0"

    status_start = max(15, bank_start + len(summary["promociones_por_banco"]) + 3)
    sheet.cell(status_start, 1, "Estado de ubicación")
    sheet.cell(status_start, 2, "Registros")
    for cell in sheet[status_start][:2]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for offset, (status, count) in enumerate(sorted(summary["estado_mapa"].items()), 1):
        sheet.cell(status_start + offset, 1, status)
        sheet.cell(status_start + offset, 2, count)
        sheet.cell(status_start + offset, 2).number_format = "#,##0"

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 4
    sheet.column_dimensions["D"].width = 24
    sheet.column_dimensions["E"].width = 16
    sheet.freeze_panes = "A5"
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            if cell.row not in {5, status_start}:
                cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")


def create_workbook(path: Path, promotions, master, locations, pending, summary):
    workbook = Workbook()
    add_summary_sheet(workbook, summary)
    master_sheet = add_data_sheet(workbook, "Base Maestra", master, MASTER_COLUMNS, "TablaBaseMaestra")
    add_data_sheet(workbook, "Promociones", promotions, PROMOTION_COLUMNS, "TablaPromociones")
    add_data_sheet(workbook, "Locales", locations, LOCATION_COLUMNS, "TablaLocales")
    pending_sheet = add_data_sheet(workbook, "Pendientes", pending, PENDING_COLUMNS, "TablaPendientes")

    status_column = MASTER_COLUMNS.index("estado_mapa") + 1
    status_letter = get_column_letter(status_column)
    master_sheet.conditional_formatting.add(
        f"{status_letter}2:{status_letter}{master_sheet.max_row}",
        FormulaRule(formula=[f'LEFT({status_letter}2,5)="LISTO"'], fill=LIGHT_GREEN),
    )
    master_sheet.conditional_formatting.add(
        f"{status_letter}2:{status_letter}{master_sheet.max_row}",
        FormulaRule(formula=[f'LEFT({status_letter}2,7)="REVISAR"'], fill=LIGHT_AMBER),
    )
    if pending:
        priority_letter = get_column_letter(PENDING_COLUMNS.index("prioridad_revision") + 1)
        pending_sheet.conditional_formatting.add(
            f"{priority_letter}2:{priority_letter}{pending_sheet.max_row}",
            FormulaRule(formula=[f'{priority_letter}2="Alta"'], fill=LIGHT_RED),
        )

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_web_records(promotions: list[dict], master: list[dict]) -> list[dict]:
    output = []
    promotion_by_id = {row.get("id_promocion", ""): row for row in promotions}
    for row in master:
        promotion = promotion_by_id.get(row["id_promocion"], {})
        precision = "none"
        if valid_coordinate(row.get("latitud"), row.get("longitud")):
            precision = "exact" if row.get("confianza_ubicacion") == "Alta" else "geocoded"
        output.append({
            "id": row["id_registro"], "promotionId": row["id_promocion"], "locationId": row["id_local"],
            "bank": row["banco"], "type": row["tipo_beneficio"], "name": row["nombre_promocion"],
            "merchant": row["comercio"], "detail": row["descripcion"], "benefit": row["beneficio_texto"],
            "discount": row["porcentaje_descuento"], "category": row["categoria"], "subcategory": row["subcategoria"],
            "sourceCategory": promotion.get("categoria_original", ""),
            "district": row["distrito"], "province": row["provincia"], "department": row["departamento"],
            "location": row["direccion_limpia"] or row["nombre_local"], "latitude": row["latitud"], "longitude": row["longitud"],
            "coordinatePrecision": precision, "mapStatus": row["estado_mapa"], "startDate": str(row["fecha_inicio"] or ""),
            "endDate": str(row["fecha_fin"] or ""), "status": row["estado_vigencia"], "paymentMethod": row["medio_pago"],
            "restrictions": row["restricciones"], "terms": row["terminos_condiciones"],
            "link": row["url_promocion"], "image": row["url_imagen"], "sourceFile": row["archivo_origen"],
        })
    return output


def build_geojson(master: list[dict]) -> dict:
    features = []
    for row in master:
        if not valid_coordinate(row.get("latitud"), row.get("longitud")):
            continue
        features.append({
            "type": "Feature", "id": row["id_registro"],
            "geometry": {"type": "Point", "coordinates": [row["longitud"], row["latitud"]]},
            "properties": {
                "id_promocion": row["id_promocion"], "id_local": row["id_local"], "banco": row["banco"],
                "comercio": row["comercio"], "promocion": row["nombre_promocion"], "categoria": row["categoria"],
                "subcategoria": row["subcategoria"], "distrito": row["distrito"], "direccion": row["direccion_limpia"],
                "beneficio": row["beneficio_texto"], "estado": row["estado_vigencia"], "imagen": row["url_imagen"],
                "url": row["url_promocion"], "confianza": row["confianza_ubicacion"],
            },
        })
    return {"type": "FeatureCollection", "features": features}


def main() -> int:
    parser = argparse.ArgumentParser(description="Genera la base maestra, el catálogo web y el mapa de promociones.")
    parser.add_argument("--entrada", default="salida", help="Carpeta que contiene los JSON procesados.")
    parser.add_argument("--salida", default="salida_final", help="Carpeta para la base maestra y archivos web.")
    args = parser.parse_args()
    input_dir, output_dir = Path(args.entrada), Path(args.salida)
    promotions, master, locations, pending, summary = build_datasets(input_dir)

    excel_path = output_dir / "Base_Maestra_Promociones.xlsx"
    create_workbook(excel_path, promotions, master, locations, pending, summary)
    save_json(output_dir / "promociones_web.json", build_web_records(promotions, master), compact=True)
    save_json(output_dir / "locales_mapa.geojson", build_geojson(master), compact=True)
    save_json(output_dir / "resumen_maestro.json", summary)

    print(f"Promociones únicas: {summary['total_promociones']:,}")
    print(f"Registros en base maestra: {summary['total_registros_maestros']:,}")
    print(f"Puntos para el mapa: {summary['locales_listos_mapa']:,}")
    print(f"Direcciones pendientes: {summary['locales_pendientes']:,}")
    print(f"Salida: {output_dir.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
