from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRAPERS = ROOT / "scripts" / "scrapers"
OUTPUT = ROOT / "data" / "fuentes"


def valid_excel(path: Path) -> bool:
    if not path.exists() or path.stat().st_size < 5_000:
        return False
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return any(sheet.max_row > 1 for sheet in workbook.worksheets)
    finally:
        workbook.close()


def run(name: str, command: list[str], target: str) -> bool:
    destination = OUTPUT / target
    temporary = OUTPUT / f".{destination.stem}.nuevo.xlsx"
    temporary.unlink(missing_ok=True)
    expanded = [str(temporary) if value == "{output}" else value for value in command]
    print(f"\n===== {name} =====", flush=True)
    result = subprocess.run([sys.executable, *expanded], cwd=ROOT, check=False)
    if result.returncode == 0 and valid_excel(temporary):
        temporary.replace(destination)
        print(f"{name}: actualizado correctamente.", flush=True)
        return True
    temporary.unlink(missing_ok=True)
    print(f"AVISO: {name} falló; se conserva el último archivo válido.", flush=True)
    return False


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    jobs = [
        ("Banco Falabella", [str(SCRAPERS / "falabella.py"), "--output", "{output}"], "promociones_banco_falabella.xlsx"),
        ("Banco Ripley", [str(SCRAPERS / "ripley.py"), "--output", "{output}"], "promociones_banco_ripley.xlsx"),
        ("Interbank", [str(SCRAPERS / "interbank.py"), "--output", "{output}"], "promociones_interbank.xlsx"),
        ("Tarjeta Cencosud", [str(SCRAPERS / "cencosud.py"), "--output", "{output}"], "Cencosud_Promociones.xlsx"),
        ("SIP", [str(SCRAPERS / "sip.py"), "--salida", "{output}"], "SIP_Promociones.xlsx"),
    ]
    results = [run(*job) for job in jobs]
    banbif_mode = "dni" if os.getenv("CLUBHOLA_DNI", "").strip() else "public"
    results.append(run("BanBif", [str(SCRAPERS / "banbif.py"), "--modo", banbif_mode, "--output", "{output}"], "BanBif_ClubHola_Promociones.xlsx"))
    successes = sum(results)
    if successes < 4:
        print(f"ERROR: solo {successes} fuentes se actualizaron correctamente.", file=sys.stderr)
        return 2
    print(f"\nFuentes actualizadas: {successes}/{len(results)}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
